from pathlib import Path
import sys
import os

ISOLATE_ROOT = Path(os.environ.get("ISOLATE_WORKSPACE", Path(__file__).resolve().parent))

"""
Growth rate r_max via sliding window on ln(OD).
Style: R theme_pub — jitter + mean±SEM + paired significance brackets.
Outputs:
  growth_rate_r_3temp.pdf/png   — 20/30/40°C + dual sig
  growth_rate_r_30v40.pdf/png   — 30 vs 40 only + sig
Exclude species10.
"""
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from scipy import stats

XLSX = str(ISOLATE_ROOT / 'Summary.xlsx')
OUTDIR = str(ISOLATE_ROOT / 'figures/growth_rate')
os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
temps = [20, 30, 40]
EXCLUDE = {'species10'}

# ── R theme_pub parameters ────────────────────────────────────
FONT_FAMILY = 'Arial'
FONT_AX     = 8
BORDER_SIZE = 0.6
TICK_SIZE   = 0.20
TICK_LEN    = 1.0

DOT_SIZE    = 4.0
JIT_SIZE    = 12
JIT_ALPHA   = 0.45
JIT_WIDTH   = 0.08
ERR_LW      = 0.45
ERR_W       = 0.10
MEAN_STROKE = 0.7

COL_TEMP = {20: '#85C1E9', 30: '#2E6DA4', 40: '#C0392B'}

plt.rcParams.update({
    'font.family': FONT_FAMILY,
    'font.size': FONT_AX,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'axes.facecolor': 'white',
    'axes.edgecolor': 'black',
    'axes.linewidth': BORDER_SIZE,
    'axes.grid': False,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
    'xtick.major.size': TICK_LEN,
    'ytick.major.size': TICK_LEN,
    'xtick.major.width': TICK_SIZE,
    'ytick.major.width': TICK_SIZE,
    'xtick.top': True,
    'ytick.right': True,
    'xtick.labeltop': False,
    'ytick.labelright': False,
    'figure.facecolor': 'white',
})

# ── Time points (0-24h) ──────────────────────────────────────
CYCLE_LENGTH = 24
time_header_ws = wb['20OD']
time_points = []
for cell in time_header_ws[1][2:]:
    if cell.value is not None and str(cell.value).strip() != '':
        try:
            h = float(cell.value)
            if h > CYCLE_LENGTH:
                break
            time_points.append(h)
        except:
            break
time_points = np.array(time_points)

# ── Read OD data ──────────────────────────────────────────────
def read_od_sheet(name):
    ws = wb[name]
    data = {}
    ncols = len(time_points)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            break
        sid = str(row[0]).strip()
        if sid in EXCLUDE:
            continue
        vals = []
        for v in row[2:2 + ncols]:
            if v is None or str(v).strip() == '':
                vals.append(np.nan)
            else:
                try:
                    vals.append(float(v))
                except:
                    vals.append(np.nan)
        data[sid] = vals
    return data

od_data = {t: read_od_sheet(f'{t}OD') for t in temps}

# ── Sliding window growth rate ────────────────────────────────
WINDOW = 3
OD_BG = 0.035  # background subtraction

def calc_max_growth_rate(time_pts, od_vals):
    n = min(len(time_pts), len(od_vals))
    t_arr = np.array(time_pts[:n])
    od_arr = np.array(od_vals[:n]) - OD_BG  # subtract background
    mask = ~np.isnan(od_arr) & (od_arr > 0)
    t_valid = t_arr[mask]
    ln_od = np.log(od_arr[mask])

    if len(t_valid) < WINDOW:
        return np.nan

    best_r = -np.inf
    for i in range(len(t_valid) - WINDOW + 1):
        t_w = t_valid[i:i + WINDOW]
        y_w = ln_od[i:i + WINDOW]
        coeffs = np.polyfit(t_w, y_w, 1)
        r = coeffs[0]
        if r > best_r:
            best_r = r
    return best_r

# Compute r for all species at all temps
r_data = {t: {} for t in temps}
for t in temps:
    for sid, od_vals in od_data[t].items():
        r = calc_max_growth_rate(time_points, od_vals)
        r_data[t][sid] = r

# Print summary
print(f'{"Species":<20s} {"20°C":>8s} {"30°C":>8s} {"40°C":>8s}')
print('-' * 48)
all_sids = sorted(r_data[20].keys())
for sid in all_sids:
    vals = [r_data[t].get(sid, np.nan) for t in temps]
    print(f'{sid:<20s} {vals[0]:>8.4f} {vals[1]:>8.4f} {vals[2]:>8.4f}')

# ── Paired significance ──────────────────────────────────────
def paired_sig(r_dict, t1, t2, one_sided_greater=False):
    common = sorted(set(r_dict[t1].keys()) & set(r_dict[t2].keys()))
    v1, v2 = [], []
    for sid in common:
        a, b = r_dict[t1][sid], r_dict[t2][sid]
        if np.isnan(a) or np.isnan(b):
            continue
        v1.append(a)
        v2.append(b)
    v1, v2 = np.array(v1), np.array(v2)
    diff = v1 - v2
    n = len(diff)
    print(f'  n={n}, mean {t1}°C={v1.mean():.4f}, {t2}°C={v2.mean():.4f}')

    sw_p = stats.shapiro(diff).pvalue if n >= 3 else 0
    normal = sw_p > 0.05
    print(f'  Shapiro diff: p={sw_p:.4f} → {"normal" if normal else "non-normal"}')

    if one_sided_greater:
        if normal:
            res = stats.ttest_rel(v2, v1)
            p = res.pvalue / 2
            if (v2 - v1).mean() < 0:
                p = 1 - p
            method = 'Paired t (one-sided)'
        else:
            res = stats.wilcoxon(v2, v1, alternative='greater')
            p = res.pvalue
            method = 'Wilcoxon (one-sided)'
    else:
        if normal:
            res = stats.ttest_rel(v1, v2)
            p = res.pvalue
            method = 'Paired t'
        else:
            res = stats.wilcoxon(v1, v2)
            p = res.pvalue
            method = 'Wilcoxon'

    print(f'  {method}: p={p:.4f}')
    if p > 0.05: label = 'ns'
    elif p > 0.01: label = '*'
    elif p > 0.001: label = '**'
    else: label = '***'
    print(f'  → {label}')
    return label

# ── Drawing helpers ───────────────────────────────────────────
def add_sig_bracket(ax, label, x1, x2, y_bar, tick_h):
    ax.plot([x1, x1, x2, x2],
            [y_bar - tick_h, y_bar, y_bar, y_bar - tick_h],
            color='black', linewidth=0.3, clip_on=False)
    ax.text((x1 + x2) / 2, y_bar + tick_h * 0.5, label,
            ha='center', va='bottom', fontsize=FONT_AX,
            fontfamily=FONT_FAMILY)

def plot_r(temp_list, color_map, filename_base, x_limits, x_breaks,
           x_labels, sig_pairs, w_mm=35, h_mm=45):
    fig, ax = plt.subplots(figsize=(w_mm / 25.4, h_mm / 25.4))

    vals_per_x = {}
    for k, t in enumerate(temp_list):
        x_pos = k + 1
        raw = [r_data[t][sid] for sid in all_sids
               if sid in r_data[t] and not np.isnan(r_data[t][sid])]
        vals_per_x[x_pos] = (t, raw)

    rng = np.random.RandomState(42)
    for x_pos, (t, vals) in vals_per_x.items():
        jx = x_pos + rng.uniform(-JIT_WIDTH, JIT_WIDTH, len(vals))
        ax.scatter(jx, vals, s=JIT_SIZE, color=color_map[t],
                   alpha=JIT_ALPHA, edgecolors='none', zorder=2)

    for x_pos, (t, vals) in vals_per_x.items():
        arr = np.array(vals)
        mu = np.mean(arr)
        sem = np.std(arr, ddof=1) / np.sqrt(len(arr)) if len(arr) > 1 else 0
        ax.errorbar(x_pos, mu, yerr=sem, fmt='none',
                    ecolor=color_map[t], elinewidth=ERR_LW,
                    capsize=ERR_W * 25, capthick=ERR_LW, zorder=3)
        ax.plot(x_pos, mu, 'o', markersize=DOT_SIZE,
                markerfacecolor='white', markeredgecolor=color_map[t],
                markeredgewidth=MEAN_STROKE, zorder=4)

    Y_MAX = 1.5
    ax.set_xlim(x_limits)
    ax.set_xticks(x_breaks)
    ax.set_xticklabels(x_labels, fontsize=FONT_AX)
    ax.set_ylim(0, Y_MAX)
    ax.set_yticks([0, 0.5, 1.0, 1.5])
    ax.set_xlabel('')
    ax.set_ylabel('')
    ax.set_title('')

    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)
    ax.tick_params(axis='both', which='both', direction='in',
                   top=True, right=True,
                   labeltop=False, labelright=False,
                   length=TICK_LEN, width=TICK_SIZE, pad=3)

    for i, (t1, t2, x1, x2, osg) in enumerate(sig_pairs):
        print(f'\nSig: {t1}°C vs {t2}°C')
        label = paired_sig(r_data, t1, t2, one_sided_greater=osg)
        y_bar = Y_MAX * (0.82 + i * 0.06)
        tick_h = Y_MAX * 0.025
        add_sig_bracket(ax, label, x1, x2, y_bar, tick_h)

    plt.tight_layout(pad=0.3)
    for ext in ['pdf', 'png']:
        outpath = os.path.join(OUTDIR, f'{filename_base}.{ext}')
        fig.savefig(outpath, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f'\nSaved: {filename_base}.pdf/.png')

# ══════════════════════════════════════════════════════════════
# Plot 1: 3 temperatures + dual significance
# ══════════════════════════════════════════════════════════════
print('=' * 50)
print('Growth rate: 3 temperatures')
print('=' * 50)
plot_r(
    temp_list=[20, 30, 40],
    color_map=COL_TEMP,
    filename_base='growth_rate_r_3temp',
    x_limits=(0.6, 3.4),
    x_breaks=[1, 2, 3],
    x_labels=['20°C', '30°C', '40°C'],
    sig_pairs=[
        (20, 30, 1, 2, False),
        (30, 40, 2, 3, False),
    ],
)

# ══════════════════════════════════════════════════════════════
# Plot 2: 30 vs 40 only + significance
# ══════════════════════════════════════════════════════════════
print('\n' + '=' * 50)
print('Growth rate: 30 vs 40')
print('=' * 50)
plot_r(
    temp_list=[30, 40],
    color_map=COL_TEMP,
    filename_base='growth_rate_r_30v40',
    x_limits=(0.6, 2.4),
    x_breaks=[1, 2],
    x_labels=['30°C', '40°C'],
    sig_pairs=[
        (30, 40, 1, 2, False),
    ],
)

print('\nDone. All in figures/growth_rate/')
