from pathlib import Path
import sys
import os

ISOLATE_ROOT = Path(os.environ.get("ISOLATE_WORKSPACE", Path(__file__).resolve().parent))

import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np

PYTHON = sys.executable
XLSX = str(ISOLATE_ROOT / 'Summary.xlsx')
OUTDIR = str(ISOLATE_ROOT / 'figures')
os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)

temps = [20, 30, 40]

def read_sheet(name, ncols=None):
    ws = wb[name]
    data = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            break
        sid = row[0]
        species_name = row[1]
        end = 2 + ncols if ncols else len(row)
        values = []
        for v in row[2:end]:
            if v is None or str(v).strip() == '':
                values.append(np.nan)
            else:
                try:
                    values.append(float(v))
                except (ValueError, TypeError):
                    values.append(np.nan)
        data[sid] = (species_name, values)
    return data

CYCLE_LENGTH = 24  # only plot first cycle (0-24h)
MAX_COLS = 99  # will be set after reading header

time_header_ws = wb['20OD']
time_points = []
for cell in time_header_ws[1][2:]:
    if cell.value is not None and str(cell.value).strip() != '':
        try:
            h = float(cell.value)
            if h > CYCLE_LENGTH:
                break
            time_points.append(h)
        except (ValueError, TypeError):
            break

MAX_COLS = len(time_points)

od_data = {}
ph_data = {}
for t in temps:
    od_data[t] = read_sheet(f'{t}OD', ncols=MAX_COLS)
    ph_data[t] = read_sheet(f'{t}pH', ncols=MAX_COLS)

species_ids = list(od_data[20].keys())

all_od = []
all_ph = []
for t in temps:
    for sid in species_ids:
        all_od.extend([v for v in od_data[t][sid][1] if not np.isnan(v) and v > 0])
        all_ph.extend([v for v in ph_data[t][sid][1] if not np.isnan(v)])

od_min = 0.02
od_max = 3.0
ph_min = 2.5
ph_max = 9.5
x_max = max(time_points) + 1

OD_COLOR = '#2166AC'
PH_COLOR = '#B2182B'

plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': 10,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'axes.linewidth': 1.0,
    'xtick.major.width': 1.0,
    'ytick.major.width': 1.0,
    'xtick.major.size': 4,
    'ytick.major.size': 4,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
})

for sid in species_ids:
    species_name = od_data[20][sid][0]

    fig, axes = plt.subplots(1, 3, figsize=(12, 3.2), sharey=False)
    fig.suptitle(species_name, fontsize=13, fontweight='bold',
                 fontstyle='italic', fontfamily='Arial', y=1.02)

    for i, t in enumerate(temps):
        ax_od = axes[i]
        ax_ph = ax_od.twinx()

        n = len(time_points)
        od_vals = od_data[t][sid][1][:n]
        ph_vals = ph_data[t][sid][1][:n]

        # OD curve
        valid_od = [(tp, v) for tp, v in zip(time_points, od_vals) if not np.isnan(v) and v > 0]
        if valid_od:
            x_od, y_od = zip(*valid_od)
            ax_od.plot(x_od, y_od, 'o-', color=OD_COLOR, markersize=5,
                       linewidth=1.8, markeredgecolor='white', markeredgewidth=0.5,
                       label='OD$_{600}$', zorder=3)

        # pH curve
        valid_ph = [(tp, v) for tp, v in zip(time_points, ph_vals) if not np.isnan(v)]
        if valid_ph:
            x_ph, y_ph = zip(*valid_ph)
            ax_ph.plot(x_ph, y_ph, 's--', color=PH_COLOR, markersize=5,
                       linewidth=1.8, markeredgecolor='white', markeredgewidth=0.5,
                       label='pH', zorder=3)

        # OD axis (log)
        ax_od.set_yscale('log')
        ax_od.set_ylim(od_min, od_max)
        ax_od.set_xlim(-0.3, x_max)
        ax_od.set_xlabel('Time (h)', fontfamily='Arial', fontsize=11)
        ax_od.set_title(f'{t} °C', fontsize=11, fontfamily='Arial', pad=8)
        ax_od.yaxis.set_major_formatter(ticker.ScalarFormatter())
        ax_od.yaxis.get_major_formatter().set_scientific(False)
        ax_od.set_yticks([0.02, 0.05, 0.1, 0.2, 0.5, 1.0, 2.0])
        ax_od.get_yaxis().set_major_formatter(ticker.FuncFormatter(
            lambda val, pos: f'{val:g}'))
        ax_od.tick_params(axis='y', colors=OD_COLOR)
        ax_od.set_xticks(time_points)

        # pH axis
        ax_ph.set_ylim(ph_min, ph_max)
        ax_ph.tick_params(axis='y', colors=PH_COLOR)
        ax_ph.set_yticks(np.arange(3, 10, 1))

        # Y-axis labels only on edges
        if i == 0:
            ax_od.set_ylabel('OD$_{600}$', fontfamily='Arial', fontsize=11,
                             color=OD_COLOR)
        else:
            ax_od.set_yticklabels([])

        if i == 2:
            ax_ph.set_ylabel('pH', fontfamily='Arial', fontsize=11,
                             color=PH_COLOR, rotation=270, labelpad=15)
        else:
            ax_ph.set_yticklabels([])

        for spine in ax_od.spines.values():
            spine.set_visible(True)
            spine.set_color('black')
        for spine in ax_ph.spines.values():
            spine.set_visible(True)
            spine.set_color('black')

        # Legend only on first panel
        if i == 0:
            lines_od, labels_od = ax_od.get_legend_handles_labels()
            lines_ph, labels_ph = ax_ph.get_legend_handles_labels()
            ax_od.legend(lines_od + lines_ph, labels_od + labels_ph,
                         loc='upper left', frameon=False, fontsize=9)

    plt.tight_layout(w_pad=2.5)
    outpath = os.path.join(OUTDIR, f'{sid}_{species_name.replace(" ", "_").replace(".", "")}.png')
    fig.savefig(outpath, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f'Saved: {outpath}')

print('\nAll figures saved.')
