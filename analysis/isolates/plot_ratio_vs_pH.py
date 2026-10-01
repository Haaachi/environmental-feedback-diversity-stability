from pathlib import Path
import sys
import os

ISOLATE_ROOT = Path(os.environ.get("ISOLATE_WORKSPACE", Path(__file__).resolve().parent))

"""
Scatter: OD/log10(CFU) ratio (x) vs endpoint pH (y)
Three temperatures, three colors. theme_pub style.
Exclude species10.
"""
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

XLSX = str(ISOLATE_ROOT / 'Summary.xlsx')
OUTDIR = str(ISOLATE_ROOT / 'figures/endpoint_24h')
os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
temps = [20, 30, 40]
EXCLUDE = {'species10'}

# ── R theme_pub parameters ────────────────────────────────────
FONT_FAMILY = 'Arial'
FONT_AX     = 8
BORDER_SIZE = 0.6
TICK_SIZE   = 0.20
TICK_LEN    = 1.0
DOT_SIZE    = 18
DOT_ALPHA   = 0.75

COL_TEMP = {20: '#85C1E9', 30: '#2E6DA4', 40: '#C0392B'}

plt.rcParams.update({
    'font.family': FONT_FAMILY,
    'font.size': FONT_AX,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'axes.facecolor': 'white',
    'axes.edgecolor': 'black',
    'axes.linewidth': BORDER_SIZE,
    'axes.grid': False,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
    'xtick.major.size': TICK_LEN,
    'ytick.major.size': TICK_LEN,
    'xtick.major.width': TICK_SIZE,
    'ytick.major.width': TICK_SIZE,
    'xtick.top': True,
    'ytick.right': True,
    'xtick.labeltop': False,
    'ytick.labelright': False,
    'figure.facecolor': 'white',
})

# ── Read data ─────────────────────────────────────────────────
def find_24h_col(sname):
    ws = wb[sname]
    for idx, cell in enumerate(ws[1]):
        if cell.value is not None:
            try:
                if float(cell.value) == 24.0: return idx + 1
            except: pass
    return None

def read_24h(sname):
    col = find_24h_col(sname)
    if col is None: return {}
    ws = wb[sname]
    data = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
        if row[0].value is None: break
        sid = str(row[0].value).strip()
        if sid in EXCLUDE: continue
        v = ws.cell(row=row[0].row, column=col).value
        if v is not None:
            try: data[sid] = float(v)
            except: pass
    return data

od = {t: read_24h(f'{t}OD') for t in temps}
ph = {t: read_24h(f'{t}pH') for t in temps}
cfu_d = {t: read_24h(f'{t}CFU') for t in temps}

# ── Compute ratio and collect scatter data ────────────────────
scatter_data = {t: {'x': [], 'y': []} for t in temps}

for t in temps:
    common = set(od[t].keys()) & set(ph[t].keys()) & set(cfu_d[t].keys())
    for sid in common:
        count = cfu_d[t][sid]
        cfu_ml = count * 200.0  # count per 5uL -> CFU/mL
        if cfu_ml > 0:
            log_cfu = np.log10(cfu_ml)
            if log_cfu > 0:
                ratio = od[t][sid] / log_cfu
                scatter_data[t]['x'].append(ratio)
                scatter_data[t]['y'].append(ph[t][sid])

# ── Plot ──────────────────────────────────────────────────────
W_mm, H_mm = 45, 45
fig, ax = plt.subplots(figsize=(W_mm / 25.4, H_mm / 25.4))

for t in temps:
    ax.scatter(scatter_data[t]['x'], scatter_data[t]['y'],
               s=DOT_SIZE, color=COL_TEMP[t], alpha=DOT_ALPHA,
               edgecolors='white', linewidths=0.3, zorder=2 + temps.index(t))

ax.set_xlim(0, 0.30)
ax.set_xticks([0, 0.1, 0.2, 0.3])
ax.set_ylim(2, 10)
ax.set_yticks([3, 5, 7, 9])

ax.set_xlabel('')
ax.set_ylabel('')
ax.set_title('')

for spine in ax.spines.values():
    spine.set_visible(True)
    spine.set_color('black')
    spine.set_linewidth(BORDER_SIZE)

ax.tick_params(axis='both', which='both', direction='in',
               top=True, right=True,
               labeltop=False, labelright=False,
               length=TICK_LEN, width=TICK_SIZE, pad=3)

plt.tight_layout(pad=0.3)
for ext in ['pdf', 'png']:
    outpath = os.path.join(OUTDIR, f'ratio_vs_pH.{ext}')
    fig.savefig(outpath, dpi=300, bbox_inches='tight', facecolor='white')
plt.close(fig)
print('Saved: ratio_vs_pH.pdf/.png')
