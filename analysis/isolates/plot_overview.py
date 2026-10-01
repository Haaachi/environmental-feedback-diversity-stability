from pathlib import Path
import sys
import os

ISOLATE_ROOT = Path(os.environ.get("ISOLATE_WORKSPACE", Path(__file__).resolve().parent))

import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

XLSX = str(ISOLATE_ROOT / 'Summary.xlsx')
OUTDIR = str(ISOLATE_ROOT / 'figures')
os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
temps = [20, 30, 40]
CYCLE_MAX = 24

# --- Read time points (0-24h) for OD/pH ---
od_times = []
for cell in wb['20OD'][1][2:]:
    if cell.value is not None and str(cell.value).strip() != '':
        try:
            h = float(cell.value)
            if h > CYCLE_MAX:
                break
            od_times.append(h)
        except:
            break

# --- Read CFU time points (within 24h) ---
cfu_times_all = []
cfu_col_indices = []  # 1-based column indices
ws_cfu_ref = wb['20CFU']  # use 20CFU as reference (clean header)
for idx, cell in enumerate(ws_cfu_ref[1]):
    if cell.value is not None and str(cell.value).strip() != '':
        try:
            h = float(cell.value)
            if 0 <= h <= CYCLE_MAX:
                cfu_times_all.append(h)
                cfu_col_indices.append(idx + 1)  # 1-based
        except:
            pass

def read_od_ph(sheet_name, n_cols):
    ws = wb[sheet_name]
    all_series = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            break
        values = []
        for v in row[2:2+n_cols]:
            if v is None or str(v).strip() == '':
                values.append(np.nan)
            else:
                try:
                    values.append(float(v))
                except:
                    values.append(np.nan)
        all_series.append(values)
    return all_series

CFU_EXCLUDE = {'species10'}  # skip Staphylococcus sp.1

def read_cfu(sheet_name):
    ws = wb[sheet_name]
    all_series = []
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row_cells[0].value is None:
            break
        sid = str(row_cells[0].value).strip()
        if sid in CFU_EXCLUDE:
            continue
        values = []
        for col_idx in cfu_col_indices:
            cell = ws.cell(row=row_cells[0].row, column=col_idx)
            if cell.value is not None and str(cell.value).strip() != '':
                try:
                    d = float(cell.value)
                    # count per 5uL -> CFU/mL
                    values.append(d * 200.0)
                except:
                    values.append(np.nan)
            else:
                values.append(np.nan)
        all_series.append(values)
    return all_series

# Load data
od_data = {t: read_od_ph(f'{t}OD', len(od_times)) for t in temps}
ph_data = {t: read_od_ph(f'{t}pH', len(od_times)) for t in temps}
cfu_data = {t: read_cfu(f'{t}CFU') for t in temps}

# --- Plot style ---
plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': 11,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'axes.linewidth': 1.0,
    'xtick.major.width': 1.0,
    'ytick.major.width': 1.0,
    'xtick.major.size': 4,
    'ytick.major.size': 4,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
})

temp_colors = {20: '#85C1E9', 30: '#2E6DA4', 40: '#C0392B'}

# --- theme_pub style (matching death example exactly) ---
BORDER_SIZE = 0.6
TICK_LEN = 3.0
TICK_WIDTH = 0.5
FONT_AX = 8
LINE_WIDTH = 1.0
LINE_ALPHA = 0.35
FIG_W = 55   # mm, same as death example
FIG_H_MAIN = 50  # mm, OD/CFU panel height
FIG_H_PH = 28    # mm, pH panel height

plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': FONT_AX,
    'axes.facecolor': 'white',
    'axes.edgecolor': 'black',
    'axes.linewidth': BORDER_SIZE,
    'axes.grid': False,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
    'xtick.major.size': TICK_LEN,
    'ytick.major.size': TICK_LEN,
    'xtick.major.width': TICK_WIDTH,
    'ytick.major.width': TICK_WIDTH,
    'xtick.top': True,
    'ytick.right': True,
    'xtick.labeltop': False,
    'ytick.labelright': False,
    'xtick.color': 'black',
    'ytick.color': 'black',
    'figure.facecolor': 'white',
})

def save_panel(filename_base, plot_func, h_mm=FIG_H_MAIN):
    fig, ax = plt.subplots(figsize=(FIG_W/25.4, h_mm/25.4))
    plot_func(ax)
    ax.set_xlabel('')
    ax.set_ylabel('')
    ax.set_title('')
    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)
    ax.tick_params(axis='both', which='both', direction='in',
                   top=True, right=True,
                   labeltop=False, labelright=False,
                   length=TICK_LEN, width=TICK_WIDTH,
                   pad=2)
    ax.tick_params(axis='x', labelsize=FONT_AX, colors='black')
    ax.tick_params(axis='y', labelsize=FONT_AX, colors='black')
    leg = ax.get_legend()
    if leg:
        leg.remove()
    plt.tight_layout(pad=0.4)
    for ext in ['pdf', 'png']:
        outpath = os.path.join(OUTDIR, f'{filename_base}.{ext}')
        fig.savefig(outpath, dpi=600, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f'Saved: {filename_base}.pdf/.png')

# --- Axis settings matching death examples ---
import matplotlib.ticker as ticker
OD_YLIM = (0.02, 2.0)
OD_YTICKS = [0.02, 0.05, 0.1, 0.2, 0.5, 1.0]
PH_YLIM = (2, 10)
PH_YTICKS = [3, 5, 7, 9]

# --- Panel A: OD (log) ---
def plot_od(ax):
    for t in temps:
        for series in od_data[t]:
            valid = [(tp, v) for tp, v in zip(od_times, series)
                     if not np.isnan(v) and v > 0]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
    ax.set_yscale('log')
    ax.set_xlim(0, 24)
    ax.set_ylim(*OD_YLIM)
    ax.set_yticks(OD_YTICKS)
    ax.set_xticks([0, 6, 12, 18, 24])
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, p: f'{v:g}'))

save_panel('overview_OD', plot_od)

# --- Panel B: pH ---
def plot_ph(ax):
    for t in temps:
        for series in ph_data[t]:
            valid = [(tp, v) for tp, v in zip(od_times, series) if not np.isnan(v)]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
    ax.set_xlim(0, 24)
    ax.set_ylim(*PH_YLIM)
    ax.set_yticks(PH_YTICKS)
    ax.set_xticks([0, 6, 12, 18, 24])

save_panel('overview_pH', plot_ph, h_mm=FIG_H_PH)

# --- Panel C: CFU (log) ---
def plot_cfu(ax):
    for t in temps:
        for series in cfu_data[t]:
            valid = [(tp, v) for tp, v in zip(cfu_times_all, series) if not np.isnan(v) and v > 0]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
    ax.set_yscale('log')
    ax.set_xlim(0, 24)
    ax.set_ylim(1e4, 1e11)
    ax.set_xticks([0, 6, 12, 18, 24])

save_panel('overview_CFU', plot_cfu)

# ============================================================
# Mean ± SEM versions (with individual lines in background)
# ============================================================

MEAN_WIDTH = 1.5
MARKER_SIZE = 4
CAP_SIZE = 2
ERR_LW = 0.8

def calc_mean_sem(time_pts, all_series):
    arr = np.array(all_series)
    means, sems = [], []
    for j in range(arr.shape[1]):
        col = arr[:, j]
        valid = col[~np.isnan(col)]
        if len(valid) > 0:
            means.append(np.mean(valid))
            sems.append(np.std(valid, ddof=1) / np.sqrt(len(valid)) if len(valid) > 1 else 0)
        else:
            means.append(np.nan)
            sems.append(np.nan)
    return np.array(means), np.array(sems)

def calc_mean_sem_log(time_pts, all_series):
    arr = np.array(all_series)
    means, lo_list, hi_list = [], [], []
    for j in range(arr.shape[1]):
        col = arr[:, j]
        valid = col[(~np.isnan(col)) & (col > 0)]
        if len(valid) > 0:
            log_v = np.log10(valid)
            m = np.mean(log_v)
            se = np.std(log_v, ddof=1) / np.sqrt(len(valid)) if len(valid) > 1 else 0
            means.append(10**m)
            lo_list.append(10**(m - se))
            hi_list.append(10**(m + se))
        else:
            means.append(np.nan)
            lo_list.append(np.nan)
            hi_list.append(np.nan)
    return np.array(means), np.array(lo_list), np.array(hi_list)

# --- OD with mean±SEM (log) ---
def plot_od_mean(ax):
    import matplotlib.ticker as ticker
    for t in temps:
        for series in od_data[t]:
            valid = [(tp, v) for tp, v in zip(od_times, series)
                     if not np.isnan(v) and v > 0]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
        # mean±SEM on log scale (geometric)
        arr = np.array(od_data[t])
        means, lo_list, hi_list = [], [], []
        for j in range(arr.shape[1]):
            col = arr[:, j]
            valid_v = col[(~np.isnan(col)) & (col > 0)]
            if len(valid_v) > 0:
                log_v = np.log10(valid_v)
                m = np.mean(log_v)
                se = np.std(log_v, ddof=1) / np.sqrt(len(valid_v)) if len(valid_v) > 1 else 0
                means.append(10**m)
                lo_list.append(10**(m - se))
                hi_list.append(10**(m + se))
            else:
                means.append(np.nan)
                lo_list.append(np.nan)
                hi_list.append(np.nan)
        means = np.array(means)
        lo_list = np.array(lo_list)
        hi_list = np.array(hi_list)
        mask = ~np.isnan(means)
        tp = np.array(od_times)[mask]
        m = means[mask]
        yerr_low = m - lo_list[mask]
        yerr_high = hi_list[mask] - m
        ax.errorbar(tp, m, yerr=[yerr_low, yerr_high], fmt='-o', color=temp_colors[t],
                    linewidth=MEAN_WIDTH, markersize=MARKER_SIZE,
                    markerfacecolor='white', markeredgecolor=temp_colors[t],
                    markeredgewidth=1.2, capsize=CAP_SIZE, capthick=ERR_LW,
                    elinewidth=ERR_LW, zorder=5)
    ax.set_yscale('log')
    ax.set_xlim(0, 24)
    ax.set_ylim(*OD_YLIM)
    ax.set_yticks(OD_YTICKS)
    ax.set_xticks([0, 6, 12, 18, 24])
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, p: f'{v:g}'))

save_panel('overview_OD_mean', plot_od_mean)

# --- pH with mean±SEM ---
def plot_ph_mean(ax):
    for t in temps:
        for series in ph_data[t]:
            valid = [(tp, v) for tp, v in zip(od_times, series) if not np.isnan(v)]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
        mean, sem = calc_mean_sem(od_times, ph_data[t])
        mask = ~np.isnan(mean)
        tp = np.array(od_times)[mask]
        ax.errorbar(tp, mean[mask], yerr=sem[mask], fmt='-o', color=temp_colors[t],
                    linewidth=MEAN_WIDTH, markersize=MARKER_SIZE,
                    markerfacecolor='white', markeredgecolor=temp_colors[t],
                    markeredgewidth=1.2, capsize=CAP_SIZE, capthick=ERR_LW,
                    elinewidth=ERR_LW, zorder=5)
    ax.set_xlim(0, 24)
    ax.set_ylim(*PH_YLIM)
    ax.set_yticks(PH_YTICKS)
    ax.set_xticks([0, 6, 12, 18, 24])

save_panel('overview_pH_mean', plot_ph_mean, h_mm=FIG_H_PH)

# --- CFU with mean±SEM (geometric) ---
def plot_cfu_mean(ax):
    for t in temps:
        for series in cfu_data[t]:
            valid = [(tp, v) for tp, v in zip(cfu_times_all, series) if not np.isnan(v) and v > 0]
            if valid:
                x, y = zip(*valid)
                ax.plot(x, y, '-', color=temp_colors[t], alpha=LINE_ALPHA,
                        linewidth=LINE_WIDTH)
        mean, lo, hi = calc_mean_sem_log(cfu_times_all, cfu_data[t])
        mask = ~np.isnan(mean)
        tp = np.array(cfu_times_all)[mask]
        m = mean[mask]
        yerr_low = m - lo[mask]
        yerr_high = hi[mask] - m
        ax.errorbar(tp, m, yerr=[yerr_low, yerr_high], fmt='-o', color=temp_colors[t],
                    linewidth=MEAN_WIDTH, markersize=MARKER_SIZE,
                    markerfacecolor='white', markeredgecolor=temp_colors[t],
                    markeredgewidth=1.2, capsize=CAP_SIZE, capthick=ERR_LW,
                    elinewidth=ERR_LW, zorder=5)
    ax.set_yscale('log')
    ax.set_xlim(0, 24)
    ax.set_ylim(1e4, 1e11)
    ax.set_xticks([0, 6, 12, 18, 24])

save_panel('overview_CFU_mean', plot_cfu_mean)

# ============================================================
# Vertical histogram distributions (horizontal bars, y-axis aligned)
# — place to the right of each overview panel
# ============================================================

HIST_W = 18  # mm width (proportional to FIG_W=55)
HIST_ALPHA = 0.45

def collect_24h_by_temp(data_dict, times_list, positive_only=False):
    """Collect 24h endpoint values per temperature."""
    idx_24 = None
    for i, t in enumerate(times_list):
        if abs(t - 24.0) < 0.01:
            idx_24 = i
            break
    result = {}
    for temp in temps:
        vals = []
        if idx_24 is not None:
            for series in data_dict[temp]:
                if idx_24 < len(series):
                    v = series[idx_24]
                    if not np.isnan(v):
                        if positive_only and v <= 0:
                            continue
                        vals.append(v)
        result[temp] = vals
    return result

def save_hist(filename_base, vals_by_temp, ylim, yticks, h_mm=45,
              log_scale=False, n_bins=20, fmt_func=None, combined_color=None):
    fig, ax = plt.subplots(figsize=(HIST_W/25.4, h_mm/25.4))

    if log_scale:
        # Bin edges in log space
        bins = np.logspace(np.log10(ylim[0]), np.log10(ylim[1]), n_bins + 1)
    else:
        bins = np.linspace(ylim[0], ylim[1], n_bins + 1)

    if combined_color:
        # Pool all temps into one histogram
        all_v = []
        for t in temps:
            all_v.extend(vals_by_temp.get(t, []))
        ax.hist(all_v, bins=bins, orientation='horizontal',
                color=combined_color, alpha=0.55, edgecolor='white',
                linewidth=0.3)
    else:
        # Stacked by temperature
        data_list = [vals_by_temp.get(t, []) for t in temps]
        colors = [temp_colors[t] for t in temps]
        ax.hist(data_list, bins=bins, orientation='horizontal',
                color=colors, alpha=HIST_ALPHA, edgecolor='white',
                linewidth=0.3, stacked=True)

    if log_scale:
        ax.set_yscale('log')
    ax.set_ylim(*ylim)
    ax.set_yticks(yticks)
    if fmt_func:
        ax.yaxis.set_major_formatter(ticker.FuncFormatter(fmt_func))

    # x-axis: count — clean, no label
    ax.set_xlim(left=0)
    ax.set_xlabel('')
    ax.set_ylabel('')
    ax.set_title('')
    # Hide y tick labels (they align with the main panel)
    ax.set_yticklabels([])

    for spine in ax.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)
    ax.tick_params(axis='both', which='both', direction='in',
                   top=True, right=True,
                   labeltop=False, labelright=False,
                   length=TICK_LEN, width=TICK_WIDTH, pad=2)
    ax.tick_params(axis='x', labelsize=FONT_AX, colors='black')
    ax.tick_params(axis='y', labelsize=FONT_AX, colors='black')

    plt.tight_layout(pad=0.3)
    for ext in ['pdf', 'png']:
        outpath = os.path.join(OUTDIR, f'{filename_base}.{ext}')
        fig.savefig(outpath, dpi=600, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print(f'Saved: {filename_base}.pdf/.png')

# OD distribution (24h endpoint only, stacked by temp)
od_24h = collect_24h_by_temp(od_data, od_times, positive_only=True)
save_hist('overview_OD_dist', od_24h,
          ylim=OD_YLIM, yticks=OD_YTICKS, h_mm=FIG_H_MAIN,
          log_scale=True, n_bins=12,
          fmt_func=lambda v, p: f'{v:g}')

# pH distribution (24h endpoint only)
ph_24h = collect_24h_by_temp(ph_data, od_times)
save_hist('overview_pH_dist', ph_24h,
          ylim=PH_YLIM, yticks=PH_YTICKS, h_mm=FIG_H_PH,
          log_scale=False, n_bins=12)
