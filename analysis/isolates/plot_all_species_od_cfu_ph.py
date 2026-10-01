from pathlib import Path
import sys
import os

ISOLATE_ROOT = Path(os.environ.get("ISOLATE_WORKSPACE", Path(__file__).resolve().parent))

import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np

XLSX = str(ISOLATE_ROOT / 'Summary.xlsx')
OUTDIR = str(ISOLATE_ROOT / 'figures/per_species_OD_CFU_pH')
os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX, data_only=True)
temps = [20, 30, 40]

# ---- theme_pub style ----
BORDER_SIZE = 0.6
TICK_LEN = 3.0
TICK_WIDTH = 0.5
FONT_AX = 8
LINE_WIDTH = 1.8
MARKER_SIZE = 5

plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': FONT_AX,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
    'axes.facecolor': 'white',
    'axes.edgecolor': 'black',
    'axes.linewidth': BORDER_SIZE,
    'axes.grid': False,
    'xtick.direction': 'in',
    'ytick.direction': 'in',
    'xtick.major.size': TICK_LEN,
    'ytick.major.size': TICK_LEN,
    'xtick.major.width': TICK_WIDTH,
    'ytick.major.width': TICK_WIDTH,
    'xtick.top': True,
    'ytick.right': True,
    'xtick.labeltop': False,
    'ytick.labelright': False,
    'figure.facecolor': 'white',
})

OD_COLOR = '#2166AC'
CFU_COLOR = '#4DAF4A'
PH_COLOR = '#B2182B'

CYCLE_MAX = 24

# ---- Read OD/pH time points (0-24h) ----
od_times = []
for cell in wb['20OD'][1][2:]:
    if cell.value is not None and str(cell.value).strip() != '':
        try:
            h = float(cell.value)
            if h > CYCLE_MAX:
                break
            od_times.append(h)
        except:
            break

# ---- CFU time/col indices ----
def get_cfu_times_cols(sheet_name):
    ws = wb[sheet_name]
    times, cols = [], []
    for idx, cell in enumerate(ws[1]):
        if cell.value is not None and str(cell.value).strip() != '':
            try:
                h = float(cell.value)
                if 0 <= h <= CYCLE_MAX:
                    times.append(h)
                    cols.append(idx + 1)
            except:
                pass
    return times, cols

# ---- Read data for one species ----
def read_od_ph_species(sheet_name, sid):
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            break
        if str(row[0]).strip() == sid:
            vals = []
            for v in row[2:2 + len(od_times)]:
                if v is None or str(v).strip() == '':
                    vals.append(np.nan)
                else:
                    try:
                        vals.append(float(v))
                    except:
                        vals.append(np.nan)
            return vals
    return []

def read_cfu_species(sheet_name, sid):
    cfu_times, cfu_cols = get_cfu_times_cols(sheet_name)
    ws = wb[sheet_name]
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row_cells[0].value is None:
            break
        if str(row_cells[0].value).strip() == sid:
            vals = []
            for col_idx in cfu_cols:
                cell = ws.cell(row=row_cells[0].row, column=col_idx)
                if cell.value is not None and str(cell.value).strip() != '':
                    try:
                        d = float(cell.value)
                        vals.append(d * 200.0)  # count per 5uL -> CFU/mL
                    except:
                        vals.append(np.nan)
                else:
                    vals.append(np.nan)
            return cfu_times, vals
    return cfu_times, []

# ---- Get species list ----
species_list = []
ws_ref = wb['20OD']
for row in ws_ref.iter_rows(min_row=2, max_row=ws_ref.max_row, values_only=True):
    if row[0] is None:
        break
    species_list.append((str(row[0]).strip(), str(row[1]).strip()))

# ---- Plot function ----
fig_w = 55 / 25.4
fig_h_main = 50 / 25.4
fig_h_ph = 28 / 25.4

def plot_one(temp, sid, species_name):
    od_vals = read_od_ph_species(f'{temp}OD', sid)
    ph_vals = read_od_ph_species(f'{temp}pH', sid)
    cfu_times, cfu_vals = read_cfu_species(f'{temp}CFU', sid)

    tag = f'{sid}_{species_name.replace(" ", "_").replace(".", "")}_{temp}C'

    # ======== OD + CFU (dual log y-axis) ========
    fig, ax_od = plt.subplots(figsize=(fig_w, fig_h_main))
    ax_cfu = ax_od.twinx()

    # OD
    valid_od = [(t, v) for t, v in zip(od_times, od_vals)
                if not np.isnan(v) and v > 0]
    if valid_od:
        xo, yo = zip(*valid_od)
        ax_od.plot(xo, yo, 'o-', color=OD_COLOR, markersize=MARKER_SIZE,
                   linewidth=LINE_WIDTH, markeredgecolor='white',
                   markeredgewidth=0.5, zorder=3)

    ax_od.set_yscale('log')
    ax_od.set_ylim(0.02, 2.0)
    ax_od.set_xlim(0, 24)
    ax_od.set_xticks([0, 6, 12, 18, 24])
    ax_od.set_yticks([0.02, 0.05, 0.1, 0.2, 0.5, 1.0])
    ax_od.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, p: f'{v:g}'))
    ax_od.tick_params(axis='y', colors=OD_COLOR, labelsize=FONT_AX)

    # CFU
    valid_cfu = [(t, v) for t, v in zip(cfu_times, cfu_vals)
                 if not np.isnan(v) and v > 0]
    if valid_cfu:
        xc, yc = zip(*valid_cfu)
        ax_cfu.plot(xc, yc, 's--', color=CFU_COLOR, markersize=MARKER_SIZE,
                    linewidth=LINE_WIDTH, markeredgecolor='white',
                    markeredgewidth=0.5, zorder=3)

    ax_cfu.set_yscale('log')
    ax_cfu.set_ylim(1e3, 1e10)
    ax_cfu.set_yticks([1e3, 1e5, 1e7, 1e9])
    ax_cfu.tick_params(axis='y', colors=CFU_COLOR, labelsize=FONT_AX,
                       labelright=True)

    ax_od.set_xlabel('')
    ax_od.set_ylabel('')
    ax_od.set_title('')

    for spine in ax_od.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)
    for spine in ax_cfu.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)

    ax_od.tick_params(axis='both', direction='in', top=True, right=True,
                      labeltop=False, labelright=False,
                      length=TICK_LEN, width=TICK_WIDTH, pad=2)
    ax_od.tick_params(axis='x', labelsize=FONT_AX)
    ax_cfu.tick_params(axis='both', direction='in', top=True, right=True,
                       labeltop=False, length=TICK_LEN, width=TICK_WIDTH, pad=2)

    plt.tight_layout(pad=0.4)
    outpath = os.path.join(OUTDIR, f'{tag}_OD_CFU.png')
    fig.savefig(outpath, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig)

    # ======== pH (shorter) ========
    fig2, ax_ph = plt.subplots(figsize=(fig_w, fig_h_ph))

    valid_ph = [(t, v) for t, v in zip(od_times, ph_vals)
                if not np.isnan(v)]
    if valid_ph:
        xp, yp = zip(*valid_ph)
        ax_ph.plot(xp, yp, 's-', color=PH_COLOR, markersize=MARKER_SIZE,
                   linewidth=LINE_WIDTH, markeredgecolor='white',
                   markeredgewidth=0.5, zorder=3)

    ax_ph.set_xlim(0, 24)
    ax_ph.set_ylim(2, 10)
    ax_ph.set_xticks([0, 6, 12, 18, 24])
    ax_ph.set_yticks([3, 5, 7, 9])

    ax_ph.set_xlabel('')
    ax_ph.set_ylabel('')
    ax_ph.set_title('')

    for spine in ax_ph.spines.values():
        spine.set_visible(True)
        spine.set_color('black')
        spine.set_linewidth(BORDER_SIZE)

    ax_ph.tick_params(axis='both', which='both', direction='in',
                      top=True, right=True,
                      labeltop=False, labelright=False,
                      length=TICK_LEN, width=TICK_WIDTH, pad=2)
    ax_ph.tick_params(axis='x', labelsize=FONT_AX)
    ax_ph.tick_params(axis='y', labelsize=FONT_AX, colors=PH_COLOR)

    plt.tight_layout(pad=0.4)
    outpath2 = os.path.join(OUTDIR, f'{tag}_pH.png')
    fig2.savefig(outpath2, dpi=300, bbox_inches='tight', facecolor='white')
    plt.close(fig2)

    return tag

# ---- Generate all ----
count = 0
for sid, species_name in species_list:
    for t in temps:
        tag = plot_one(t, sid, species_name)
        count += 1
        print(f'[{count}] {tag}')

print(f'\nDone. {count} figure pairs saved to {OUTDIR}')
