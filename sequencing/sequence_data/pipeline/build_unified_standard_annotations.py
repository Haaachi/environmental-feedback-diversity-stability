"""Build unified annotation tables from the standard per-experiment outputs."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from annotate_species_local_blast import annotation_paths, standard_annotation_dir
from synthetic_community_pipeline import EXPERIMENTS, WORKSPACE_ROOT, output_paths


OUT_DIR = WORKSPACE_ROOT / "species_annotations" / "unified_standard"


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def parse_taxon_levels(taxon: str) -> dict[str, str]:
    levels: dict[str, str] = {}
    for part in str(taxon or "").split(";"):
        part = part.strip()
        if "__" not in part:
            continue
        prefix, value = part.split("__", 1)
        value = value.strip()
        if value:
            levels[prefix] = value
    return levels


def normalize_scientific_name(name: str, genus_hint: str = "") -> str:
    text = str(name or "").strip()
    if not text or text in {"N/A", "NA", "nan", "None"}:
        return ""
    text = re.sub(r"\s+", " ", text)
    text = re.split(r"\s+(strain|subsp\.|DSM|ATCC|KCTC|LMG|JCM|NBRC|NCTC)\b", text)[0].strip()
    pieces = text.split()
    if len(pieces) >= 2 and pieces[0][0].isupper():
        return f"{pieces[0]} {pieces[1]}"
    if genus_hint and len(pieces) >= 1 and pieces[0].islower():
        return f"{genus_hint} {pieces[0]}"
    return ""


def fallback_label(levels: dict[str, str]) -> str:
    for key in ("f", "o", "c", "p"):
        if levels.get(key):
            return f"Unclassified_{levels[key]}"
    return "Unclassified"


def extract_annotation_parts(row: dict[str, str]) -> tuple[str, str, str]:
    taxon = row.get("Taxon", "") or row.get("Representative_Taxon", "")
    levels = parse_taxon_levels(taxon)
    genus = levels.get("g", "").strip()
    species_name = levels.get("s", "").strip()

    scientific = row.get("Scientific_name", "")
    if not species_name:
        species_name = normalize_scientific_name(scientific, genus)
    if species_name and not genus:
        genus = species_name.split()[0]

    genus_or_fallback = genus if genus else fallback_label(levels)
    return genus_or_fallback, species_name, levels.get("f", "")


def species_sort_key(species_id: str) -> tuple[int, str]:
    match = re.search(r"(\d+)$", species_id)
    return (int(match.group(1)) if match else 10**9, species_id)


def fallback_species_rows(experiment_key: str) -> list[dict[str, object]]:
    config = EXPERIMENTS[experiment_key]
    species_table = output_paths(config)["species_table"]
    if not species_table.exists():
        return []
    rows: list[dict[str, object]] = []
    for raw in read_tsv(species_table):
        genus, species_name, family = extract_annotation_parts(raw)
        rows.append(
            {
                "Experiment": experiment_key,
                "Species": raw.get("Species", ""),
                "Display_Taxon": raw.get("Display_Taxon", ""),
                "Feature_ID": raw.get("Representative_Original_ID", ""),
                "Representative_ASV": raw.get("Representative_ASV", ""),
                "Genus": raw.get("Genus", "") or genus,
                "Family": family,
                "Species_name": raw.get("Species_name_from_BLAST", "") or species_name,
                "Original_Taxon": raw.get("Representative_Taxon", ""),
                "Confidence": raw.get("Representative_Taxonomy_confidence", ""),
                "Identity_pct": "",
                "Coverage_pct": "",
                "Top_accession": "",
                "Top_hit_description": "",
                "Annotation_source": "standard_species_table_fallback",
            }
        )
    return rows


def annotated_species_rows(experiment_key: str) -> list[dict[str, object]]:
    config = EXPERIMENTS[experiment_key]
    mapping_path = annotation_paths(config, standard_annotation_dir(config))["mapping"]
    if not mapping_path.exists():
        return fallback_species_rows(experiment_key)

    rows: list[dict[str, object]] = []
    for raw in read_tsv(mapping_path):
        genus, species_name, family = extract_annotation_parts(raw)
        rows.append(
            {
                "Experiment": experiment_key,
                "Species": raw.get("Species", ""),
                "Display_Taxon": raw.get("Display_Taxon", ""),
                "Feature_ID": raw.get("Feature ID", ""),
                "Representative_ASV": raw.get("Representative_ASV", ""),
                "Genus": genus,
                "Family": family,
                "Species_name": species_name,
                "Original_Taxon": raw.get("Taxon", ""),
                "Confidence": raw.get("Confidence", ""),
                "Identity_pct": raw.get("Identity_pct", ""),
                "Coverage_pct": raw.get("Coverage_pct", ""),
                "Top_accession": raw.get("Top_accession", ""),
                "Top_hit_description": raw.get("Top_hit_description", ""),
                "Annotation_source": raw.get("Annotation_source", "local_blast"),
            }
        )
    return rows


def assign_unified_labels(rows: list[dict[str, object]]) -> None:
    counters: dict[tuple[str, str], int] = defaultdict(int)
    for row in rows:
        experiment = str(row["Experiment"])
        genus = str(row.get("Genus", "") or "Unclassified")
        counters[(experiment, genus)] += 1
        row["Unified_annotation"] = f"{genus} sp.{counters[(experiment, genus)]}"


def build_species_rows() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for experiment_key in EXPERIMENTS:
        rows.extend(annotated_species_rows(experiment_key))
    assign_unified_labels(rows)
    rows.sort(key=lambda row: (str(row["Experiment"]), species_sort_key(str(row["Species"]))))
    return rows


def build_community_rows(species_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    by_key = {(row["Experiment"], row["Species"]): row for row in species_rows}
    out: list[dict[str, object]] = []

    for experiment_key, config in EXPERIMENTS.items():
        community_path = output_paths(config)["community_species"]
        if not community_path.exists():
            continue
        for raw in read_tsv(community_path):
            ann = by_key.get((experiment_key, raw["Species"]), {})
            out.append(
                {
                    "Experiment": experiment_key,
                    "Community": raw.get("Community", ""),
                    "Rank": raw.get("Rank", ""),
                    "Species": raw.get("Species", ""),
                    "Unified_annotation": ann.get("Unified_annotation", ""),
                    "Display_Taxon": ann.get("Display_Taxon", ""),
                    "Genus": ann.get("Genus", ""),
                    "Species_name": ann.get("Species_name", ""),
                    "Feature_ID": raw.get("Representative_Original_ID", ann.get("Feature_ID", "")),
                    "Representative_ASV": raw.get("Representative_ASV", ann.get("Representative_ASV", "")),
                    "Max_relabund_in_community": raw.get("Max_relabund_in_community", ""),
                    "Num_samples_above_threshold": raw.get("Num_samples_above_threshold", ""),
                    "Presence_rule": raw.get("Presence_rule", ""),
                }
            )
    out.sort(
        key=lambda row: (
            str(row["Experiment"]),
            int(row["Community"]) if str(row["Community"]).isdigit() else 10**9,
            int(row["Rank"]) if str(row["Rank"]).isdigit() else 10**9,
        )
    )
    return out


def write_xlsx(path: Path, sheets: dict[str, tuple[list[dict[str, object]], list[str]]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name, (rows, fields) in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for col_idx, field in enumerate(fields, 1):
            cell = ws.cell(row=1, column=col_idx, value=field)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_idx, row in enumerate(rows, 2):
            for col_idx, field in enumerate(fields, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(field, ""))
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=field in {"Original_Taxon", "Top_hit_description"},
                )
        for col_idx, field in enumerate(fields, 1):
            width = 14
            if field in {"Original_Taxon", "Top_hit_description"}:
                width = 70
            elif field == "Feature_ID":
                width = 36
            elif field in {"Unified_annotation", "Species_name"}:
                width = 26
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    try:
        wb.save(path)
        return path
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_updated{path.suffix}")
        wb.save(fallback)
        return fallback


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    species_rows = build_species_rows()
    community_rows = build_community_rows(species_rows)

    species_fields = [
        "Experiment",
        "Species",
        "Unified_annotation",
        "Display_Taxon",
        "Genus",
        "Species_name",
        "Feature_ID",
        "Representative_ASV",
        "Family",
        "Confidence",
        "Identity_pct",
        "Coverage_pct",
        "Top_accession",
        "Original_Taxon",
        "Top_hit_description",
        "Annotation_source",
    ]
    community_fields = [
        "Experiment",
        "Community",
        "Rank",
        "Species",
        "Unified_annotation",
        "Display_Taxon",
        "Genus",
        "Species_name",
        "Feature_ID",
        "Representative_ASV",
        "Max_relabund_in_community",
        "Num_samples_above_threshold",
        "Presence_rule",
    ]

    write_tsv(OUT_DIR / "Unified_species_annotation_mapping.tsv", species_rows, species_fields)
    write_tsv(OUT_DIR / "Unified_community_species_annotations.tsv", community_rows, community_fields)
    xlsx = write_xlsx(
        OUT_DIR / "Unified_species_annotations.xlsx",
        {
            "Species_mapping": (species_rows, species_fields),
            "Community_species": (community_rows, community_fields),
        },
    )

    print(f"Wrote: {xlsx}")
    print(f"Species rows: {len(species_rows)}")
    print(f"Community rows: {len(community_rows)}")


if __name__ == "__main__":
    main()
