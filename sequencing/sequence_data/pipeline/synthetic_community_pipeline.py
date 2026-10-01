"""Standard ASV-to-synthetic-community species pipeline.

This script is the canonical shared implementation for the temperature and
mortality synthetic-community post-processing workflows. Experiment-specific
differences stay in ``EXPERIMENTS``; all table parsing, ASV grouping, species
presence filtering, QC, and Excel export logic is shared.
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]

EXCEL_SLOTS = (
    [(1, day) for day in range(1, 11)]
    + [(2, day) for day in (8, 9, 10)]
    + [(3, day) for day in (8, 9, 10)]
)
EXCEL_HEADERS = ["Taxon"] + [f"D{day}(R{rep})" for rep, day in EXCEL_SLOTS]

R_THRESHOLD = 0.95
RELABUND_MASK = 0.01
MIN_NONZERO_FOR_PRESENCE = 3
HIGH_RELABUND_FOR_PRESENCE = 0.05
SEQ_IDENTITY_FOR_MERGE = 0.995
SEQ_IDENTITY_FOR_PROFILE_MERGE = 0.99
SEQ_ASSISTED_R_THRESHOLD = 0.90
SEQ_ASSISTED_JACCARD_THRESHOLD = 0.85
SATELLITE_TOTAL_RATIO_MAX = 0.30

SAMPLE_RE_TEMPLATE = r"^{prefix}_(W\d+)_(C\d+)_(R\d+)_(D\d+)$"


@dataclass(frozen=True)
class ExperimentConfig:
    key: str
    title: str
    root: Path
    sample_prefix: str
    top_n_asv: int
    analysis_filter: Callable[[dict[str, object]], bool]
    output_dir_name: str = "species_identification_standard"
    taxonomy_file: Path | None = None
    readme_notes: tuple[str, ...] = field(default_factory=tuple)

    @property
    def input_table(self) -> Path:
        return self.root / "exported" / "feature-table.tsv"

    @property
    def fasta_file(self) -> Path:
        return self.root / "exported" / "dna-sequences.fasta"

    @property
    def output_dir(self) -> Path:
        return self.root / self.output_dir_name

    @property
    def label(self) -> str:
        return self.title


def all_samples(_: dict[str, object]) -> bool:
    return True


TEMPERATURE_W5_VALID_COMMUNITIES = {3, 7, 12}
TEMPERATURE_W5_COLLAPSED_SLOTS = {
    (7, 1, 7),
    (7, 2, 7),
    (7, 3, 5),
    (7, 3, 6),
    (7, 3, 7),
    (12, 1, 6),
    (12, 1, 7),
    (12, 1, 8),
    (12, 3, 6),
    (12, 3, 8),
}


def temperature_analysis_samples(meta: dict[str, object]) -> bool:
    if meta["W"] != "W5":
        return True
    cnum = int(meta["C_num"])
    if cnum not in TEMPERATURE_W5_VALID_COMMUNITIES:
        return False
    slot = (cnum, int(meta["R_num"]), int(meta["D_num"]))
    return slot not in TEMPERATURE_W5_COLLAPSED_SLOTS


def mortality_analysis_samples(meta: dict[str, object]) -> bool:
    return meta["R"] == "R1" and 1 <= int(meta["D_num"]) <= 6


EXPERIMENTS: dict[str, ExperimentConfig] = {
    "temperature": ExperimentConfig(
        key="temperature",
        title="Temperature",
        root=WORKSPACE_ROOT / "16s_project",
        sample_prefix="temperature",
        top_n_asv=45,
        analysis_filter=temperature_analysis_samples,
        taxonomy_file=WORKSPACE_ROOT
        / "16s_project"
        / "exported"
        / "local_blast_taxonomy"
        / "taxonomy.tsv",
        readme_notes=(
            "Analysis/modeling samples: temperature W1-W4 plus non-collapsed W5 samples from communities 3, 7, and 12. Other W5 communities, and OD-collapsed samples within C7/C12, are excluded from TopN ranking, ASV merging, presence calling, and QC because their relative-abundance profiles are noise-dominated.",
            "In the current feature table, W5 R2/R3 D8-D9 are measured only for C3, C7, and C12.",
            "Temperature standard TOP_N_ASV is set to 45 as a conservative interpretation point; W5 communities 3, 7, and 12 are treated as meaningful sequenced communities only at OD-supported sample points. W5 C7 R1/R2 D7, C7 R3 D5-D7, C12 R1 D6-D8, and C12 R3 D6/D8 are excluded as OD-collapsed sample points.",
            "Taxonomy source, when available: exported/local_blast_taxonomy/taxonomy.tsv from the all-ASV BLAST step.",
        ),
    ),
    "mortality": ExperimentConfig(
        key="mortality",
        title="Mortality",
        root=WORKSPACE_ROOT / "Mortality_analysis",
        sample_prefix="mortality",
        top_n_asv=45,
        analysis_filter=mortality_analysis_samples,
        taxonomy_file=WORKSPACE_ROOT / "Mortality_analysis" / "exported" / "taxonomy.tsv",
        readme_notes=(
            "Analysis samples: mortality_*_R1_D1-D6.",
            "Missing plotting slots D7-D10(R1) and D8-D10(R2/R3) are filled with 1 for plotting compatibility.",
            "QIIME taxonomy is carried through as a fallback; the standard local BLAST annotation step can overwrite/extend it.",
        ),
    ),
}


def output_paths(config: ExperimentConfig) -> dict[str, Path]:
    out = config.output_dir
    label = config.label
    return {
        "xlsx": out / f"{label}_Unified_final.xlsx",
        "asv_mapping": out / f"{label}_asv_mapping.tsv",
        "species_table": out / f"{label}_species_table.tsv",
        "species_fasta": out / f"{label}_species.fasta",
        "community_species": out / f"{label}_community_species.tsv",
        "species_count": out / f"{label}_species_count.tsv",
        "retention_qc": out / f"{label}_retention_qc.tsv",
        "pair_diagnostics": out / f"{label}_asv_pair_diagnostics.tsv",
    }


def read_feature_table(path: Path) -> pd.DataFrame:
    with path.open("r", encoding="utf-8") as handle:
        skiprows = 1 if handle.readline().startswith("# Constructed") else 0
    df = pd.read_csv(path, sep="\t", index_col=0, skiprows=skiprows)
    df.index.name = "Feature_ID"
    return df.apply(pd.to_numeric, errors="coerce").fillna(0)


def parse_sample_name(sample: str, prefix: str) -> dict[str, object]:
    pattern = re.compile(SAMPLE_RE_TEMPLATE.format(prefix=re.escape(prefix)))
    match = pattern.match(sample)
    if not match:
        raise ValueError(f"Unexpected sample name for {prefix}: {sample}")
    w, c, r, d = match.groups()
    return {
        "W": w,
        "C": c,
        "R": r,
        "D": d,
        "W_num": int(w[1:]),
        "C_num": int(c[1:]),
        "R_num": int(r[1:]),
        "D_num": int(d[1:]),
    }


def read_fasta(path: Path) -> dict[str, str]:
    seqs: dict[str, str] = {}
    name = None
    parts: list[str] = []
    with path.open("r", encoding="utf-8") as handle:
        for raw in handle:
            line = raw.strip()
            if not line:
                continue
            if line.startswith(">"):
                if name is not None:
                    seqs[name] = "".join(parts)
                name = line[1:].split()[0]
                parts = []
            else:
                parts.append(line)
    if name is not None:
        seqs[name] = "".join(parts)
    return seqs


def read_taxonomy(path: Path | None) -> tuple[dict[str, str], dict[str, str]]:
    if path is None or not path.exists():
        return {}, {}
    taxonomy = pd.read_csv(path, sep="\t", keep_default_na=False)
    tax_by_id = dict(zip(taxonomy.get("Feature ID", []), taxonomy.get("Taxon", [])))
    conf_by_id = dict(zip(taxonomy.get("Feature ID", []), taxonomy.get("Confidence", [])))
    return tax_by_id, conf_by_id


def parse_taxon_levels(taxon: str) -> dict[str, str]:
    levels: dict[str, str] = {}
    for part in str(taxon or "").split(";"):
        part = part.strip()
        if "__" not in part:
            continue
        prefix, value = part.split("__", 1)
        value = value.strip()
        if value:
            levels[prefix] = value
    return levels


def genus_or_fallback(taxon: str) -> str:
    levels = parse_taxon_levels(taxon)
    if levels.get("g"):
        return levels["g"]
    for key in ("f", "o", "c", "p"):
        if levels.get(key):
            return f"Unclassified_{levels[key]}"
    return "Unclassified"


def species_name_from_taxon(taxon: str) -> str:
    return parse_taxon_levels(taxon).get("s", "")


def informative_genus(genus: str) -> bool:
    text = str(genus or "")
    return bool(text) and not text.startswith("Unclassified")


def assign_display_taxa(species_table_rows: list[dict[str, object]]) -> dict[str, str]:
    counters: dict[str, int] = defaultdict(int)
    display_by_species: dict[str, str] = {}
    for row in species_table_rows:
        genus = str(row.get("Genus") or "Unclassified")
        counters[genus] += 1
        display = f"{genus} sp.{counters[genus]}"
        row["Display_Taxon"] = display
        display_by_species[str(row["Species"])] = display
    return display_by_species


def levenshtein_distance(a: str, b: str) -> int:
    if a == b:
        return 0
    if len(a) < len(b):
        a, b = b, a
    previous = list(range(len(b) + 1))
    for i, char_a in enumerate(a, 1):
        current = [i]
        for j, char_b in enumerate(b, 1):
            current.append(
                min(
                    previous[j] + 1,
                    current[j - 1] + 1,
                    previous[j - 1] + (char_a != char_b),
                )
            )
        previous = current
    return previous[-1]


def sequence_identity(seq_a: str, seq_b: str) -> float:
    if not seq_a or not seq_b:
        return 0.0
    if len(seq_a) == len(seq_b):
        matches = sum(base_a == base_b for base_a, base_b in zip(seq_a, seq_b))
        return matches / len(seq_a) if seq_a else 0.0
    distance = levenshtein_distance(seq_a, seq_b)
    return 1.0 - (distance / max(len(seq_a), len(seq_b)))


def jaccard_similarity(mask_a: np.ndarray, mask_b: np.ndarray) -> float:
    union = np.logical_or(mask_a, mask_b).sum()
    if union == 0:
        return 0.0
    return float(np.logical_and(mask_a, mask_b).sum() / union)


def build_pair_diagnostics(
    top_asv_list: list[str],
    corr: np.ndarray,
    mat_rel: pd.DataFrame,
    total_reads_top: pd.Series,
    short_to_orig: dict[str, str],
    seqs_by_original_id: dict[str, str],
    tax_by_id: dict[str, str],
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    presence = mat_rel > RELABUND_MASK

    for idx_i in range(len(top_asv_list)):
        asv_i = top_asv_list[idx_i]
        orig_i = short_to_orig[asv_i]
        tax_i = tax_by_id.get(orig_i, "")
        genus_i = genus_or_fallback(tax_i)
        seq_i = seqs_by_original_id.get(orig_i, "")
        reads_i = float(total_reads_top[asv_i])
        mask_i = presence.loc[asv_i].values

        for idx_j in range(idx_i + 1, len(top_asv_list)):
            asv_j = top_asv_list[idx_j]
            orig_j = short_to_orig[asv_j]
            tax_j = tax_by_id.get(orig_j, "")
            genus_j = genus_or_fallback(tax_j)
            seq_j = seqs_by_original_id.get(orig_j, "")
            reads_j = float(total_reads_top[asv_j])
            seq_id = sequence_identity(seq_i, seq_j)
            pearson_r = float(corr[idx_i, idx_j])
            jaccard = jaccard_similarity(mask_i, presence.loc[asv_j].values)
            minor_major_ratio = min(reads_i, reads_j) / max(reads_i, reads_j) if max(reads_i, reads_j) else 0.0
            same_genus = genus_i == genus_j and informative_genus(genus_i)
            profile_merge_edge = (
                pearson_r > R_THRESHOLD
                and same_genus
                and (
                    seq_id >= SEQ_IDENTITY_FOR_PROFILE_MERGE
                    or minor_major_ratio <= SATELLITE_TOTAL_RATIO_MAX
                )
            )
            sequence_merge_edge = (
                same_genus
                and seq_id >= SEQ_IDENTITY_FOR_MERGE
                and pearson_r >= SEQ_ASSISTED_R_THRESHOLD
                and jaccard >= SEQ_ASSISTED_JACCARD_THRESHOLD
                and minor_major_ratio <= SATELLITE_TOTAL_RATIO_MAX
            )
            if sequence_merge_edge:
                recommendation = "merge_sequence_supported_satellite"
            elif profile_merge_edge:
                recommendation = "merge_profile_supported"
            elif same_genus and seq_id >= SEQ_IDENTITY_FOR_MERGE:
                recommendation = "review_high_sequence_similarity"
            elif pearson_r > R_THRESHOLD:
                recommendation = "review_high_correlation_not_auto_merged"
            else:
                recommendation = "do_not_merge"

            rows.append(
                {
                    "ASV_i": asv_i,
                    "ASV_j": asv_j,
                    "Original_ID_i": orig_i,
                    "Original_ID_j": orig_j,
                    "Genus_i": genus_i,
                    "Genus_j": genus_j,
                    "Same_informative_genus": same_genus,
                    "Sequence_identity": round(seq_id, 6),
                    "Pearson_r": round(pearson_r, 6),
                    "Presence_jaccard": round(jaccard, 6),
                    "Total_reads_i": int(reads_i),
                    "Total_reads_j": int(reads_j),
                    "Minor_major_total_ratio": round(minor_major_ratio, 6),
                    "Profile_merge_edge": profile_merge_edge,
                    "Sequence_merge_edge": sequence_merge_edge,
                    "Merge_edge": profile_merge_edge or sequence_merge_edge,
                    "Recommendation": recommendation,
                    "Taxon_i": tax_i,
                    "Taxon_j": tax_j,
                }
            )

    return pd.DataFrame(rows)


def connected_components(edge_i: np.ndarray, edge_j: np.ndarray, n_nodes: int) -> list[list[int]]:
    parent = list(range(n_nodes))

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a: int, b: int) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    for i, j in zip(edge_i, edge_j):
        union(int(i), int(j))

    comp: dict[int, list[int]] = defaultdict(list)
    for idx in range(n_nodes):
        comp[find(idx)].append(idx)
    return list(comp.values())


def python_value(value: object) -> object:
    if isinstance(value, np.generic):
        return value.item()
    return value


def write_table_sheet(ws, table: pd.DataFrame) -> None:
    for col_idx, col in enumerate(table.columns, 1):
        ws.cell(row=1, column=col_idx, value=col)
    for row_idx, record in enumerate(table.itertuples(index=False), 2):
        for col_idx, value in enumerate(record, 1):
            ws.cell(row=row_idx, column=col_idx, value=python_value(value))


def style_excel(wb: Workbook) -> None:
    title_fill = PatternFill("solid", fgColor="1F3864")
    title_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(color="FFFFFF", bold=True)
    zero_font = Font(color="BBBBBB")
    fill_font = Font(color="999999")

    for ws in wb.worksheets:
        ws.freeze_panes = "B2" if ws.title.startswith("W") else "A2"
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        for col_idx in range(1, ws.max_column + 1):
            width = 24 if col_idx == 1 else 12
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        if not ws.title.startswith("W"):
            if ws.max_row >= 1:
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
            continue

        for row in ws.iter_rows():
            first_value = row[0].value
            if first_value and str(first_value).startswith("Community "):
                for cell in row[: len(EXCEL_HEADERS)]:
                    cell.fill = title_fill
                    cell.font = title_font
            elif first_value == "Taxon":
                for cell in row[: len(EXCEL_HEADERS)]:
                    cell.fill = header_fill
                    cell.font = header_font
            else:
                for cell in row[1 : len(EXCEL_HEADERS)]:
                    if cell.value == 0:
                        cell.font = zero_font
                    elif cell.value == 1:
                        cell.font = fill_font


def save_workbook(wb: Workbook, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_updated{path.suffix}")
        wb.save(fallback)
        return fallback


def run_pipeline(config: ExperimentConfig) -> dict[str, Path]:
    config.output_dir.mkdir(parents=True, exist_ok=True)
    paths = output_paths(config)

    print(f"[{config.title}] [1/8] Read feature table: {config.input_table}")
    df = read_feature_table(config.input_table)
    sample_prefix = f"{config.sample_prefix}_"
    community_cols = [col for col in df.columns if col.startswith(sample_prefix)]
    col_meta = {col: parse_sample_name(col, config.sample_prefix) for col in community_cols}
    analysis_cols = [col for col in community_cols if config.analysis_filter(col_meta[col])]
    if not community_cols:
        raise ValueError(f"No columns found with prefix {sample_prefix}")
    if not analysis_cols:
        raise ValueError(f"No analysis columns selected for {config.title}")
    print(f"  Raw table: {df.shape[0]} ASVs x {df.shape[1]} samples")
    print(f"  Community sample columns: {len(community_cols)}")
    print(f"  Analysis sample columns: {len(analysis_cols)}")

    print(f"[{config.title}] [2/8] Rank ASVs by analysis reads")
    total_reads_analysis = df[analysis_cols].sum(axis=1)
    total_reads_all = df[community_cols].sum(axis=1)
    order_df = pd.DataFrame(
        {
            "analysis": total_reads_analysis,
            "all": total_reads_all,
            "id": df.index,
        }
    ).sort_values(
        by=["analysis", "all", "id"],
        ascending=[False, False, True],
        kind="mergesort",
    )
    df = df.loc[order_df.index]

    original_asvs = df.index.tolist()
    asv_short = {asv: f"ASV_{idx + 1}" for idx, asv in enumerate(original_asvs)}
    short_to_orig = {short: orig for orig, short in asv_short.items()}
    df = df.rename(index=asv_short)
    df.index.name = "ASV"

    tax_by_id, conf_by_id = read_taxonomy(config.taxonomy_file)
    mapping_df = pd.DataFrame(
        {
            "Short_ID": [asv_short[asv] for asv in original_asvs],
            "Original_ID": original_asvs,
            "Taxon": [tax_by_id.get(asv, "") for asv in original_asvs],
            "Taxonomy_confidence": [conf_by_id.get(asv, "") for asv in original_asvs],
            "Total_reads_in_analysis": [int(total_reads_analysis[asv]) for asv in original_asvs],
            "Total_reads_all_samples": [int(total_reads_all[asv]) for asv in original_asvs],
        }
    )
    mapping_df.to_csv(paths["asv_mapping"], sep="\t", index=False)

    print(f"[{config.title}] [3/8] Keep top {config.top_n_asv} ASVs")
    mat_reads_analysis = df[analysis_cols].iloc[: config.top_n_asv].copy()
    mat_reads_analysis = mat_reads_analysis.loc[mat_reads_analysis.sum(axis=1) > 0]
    top_asv_list = mat_reads_analysis.index.tolist()
    if not top_asv_list:
        raise ValueError(f"No top ASVs remained for {config.title}")
    print(
        f"  Top range: {top_asv_list[0]}={int(mat_reads_analysis.iloc[0].sum())} reads, "
        f"{top_asv_list[-1]}={int(mat_reads_analysis.iloc[-1].sum())} reads"
    )

    sample_total_reads = df[analysis_cols].sum(axis=0)
    top_total_per_sample = mat_reads_analysis.sum(axis=0)
    total_reads_top = mat_reads_analysis.sum(axis=1)
    mat_rel = mat_reads_analysis.divide(top_total_per_sample.replace(0, np.nan), axis=1).fillna(0)

    print(f"[{config.title}] [4/8] Pearson r clustering, threshold > {R_THRESHOLD}")
    x = mat_rel.values.astype(float)
    std_x = x.std(axis=1, keepdims=True, ddof=0)
    z = (x - x.mean(axis=1, keepdims=True)) / np.where(std_x == 0, np.nan, std_x)
    corr = z @ z.T / x.shape[1]
    corr[np.isnan(corr)] = -np.inf
    np.fill_diagonal(corr, -np.inf)
    seqs_by_original_id = read_fasta(config.fasta_file)
    pair_diagnostics_df = build_pair_diagnostics(
        top_asv_list=top_asv_list,
        corr=corr,
        mat_rel=mat_rel,
        total_reads_top=total_reads_top,
        short_to_orig=short_to_orig,
        seqs_by_original_id=seqs_by_original_id,
        tax_by_id=tax_by_id,
    )
    pair_diagnostics_df.to_csv(paths["pair_diagnostics"], sep="\t", index=False)
    merge_pairs = pair_diagnostics_df[pair_diagnostics_df["Merge_edge"]]
    asv_position = {asv: idx for idx, asv in enumerate(top_asv_list)}
    edge_i = merge_pairs["ASV_i"].map(asv_position).to_numpy(dtype=int)
    edge_j = merge_pairs["ASV_j"].map(asv_position).to_numpy(dtype=int)
    print(f"  Pearson-only pairs above threshold: {int(np.triu(corr > R_THRESHOLD, k=1).sum())}")
    print(f"  Accepted merge edges after sequence/taxonomy checks: {len(edge_i)}")

    print(f"[{config.title}] [5/8] Build connected ASV groups")
    comp_idx = connected_components(edge_i, edge_j, len(top_asv_list))
    groups_raw = [[top_asv_list[idx] for idx in group] for group in comp_idx]
    group_representative: dict[int, str] = {}
    group_to_members: dict[int, list[str]] = {}
    for gid, members in enumerate(groups_raw):
        best = max(members, key=lambda asv: (total_reads_top[asv], -int(asv.split("_")[1])))
        group_representative[gid] = best
        group_to_members[gid] = sorted(members, key=lambda asv: int(asv.split("_")[1]))
    multi_member = sum(len(members) > 1 for members in group_to_members.values())
    print(f"  Groups: {len(group_to_members)}; multi-member groups: {multi_member}")

    print(f"[{config.title}] [6/8] Merge reads and name candidate species")
    group_reads_analysis: dict[int, pd.Series] = {}
    group_reads_all: dict[int, pd.Series] = {}
    group_total_reads_analysis: dict[int, int] = {}
    group_total_reads_all: dict[int, int] = {}
    for gid, members in group_to_members.items():
        analysis_sum = mat_reads_analysis.loc[members].sum(axis=0)
        all_sum = df.loc[members, community_cols].sum(axis=0)
        group_reads_analysis[gid] = analysis_sum
        group_reads_all[gid] = all_sum
        group_total_reads_analysis[gid] = int(analysis_sum.sum())
        group_total_reads_all[gid] = int(all_sum.sum())

    ordered_gids = sorted(
        group_to_members,
        key=lambda gid: (
            -group_total_reads_analysis[gid],
            int(group_representative[gid].split("_")[1]),
        ),
    )
    candidate_name = {gid: f"candidate{idx + 1}" for idx, gid in enumerate(ordered_gids)}
    candidate_to_gid = {species: gid for gid, species in candidate_name.items()}
    candidate_names = [candidate_name[gid] for gid in ordered_gids]
    candidate_reads_df = pd.DataFrame(
        [group_reads_analysis[gid].values for gid in ordered_gids],
        index=candidate_names,
        columns=analysis_cols,
    )
    candidate_rel_df = candidate_reads_df.divide(
        top_total_per_sample.replace(0, np.nan), axis=1
    ).fillna(0)
    print(f"  Candidate groups: {len(ordered_gids)}")

    print(f"[{config.title}] [7/8] Detect species present in each community")
    print(
        "  Presence rule: "
        f"(samples with rel_abund > {RELABUND_MASK} >= {MIN_NONZERO_FOR_PRESENCE}) "
        f"OR (max rel_abund >= {HIGH_RELABUND_FOR_PRESENCE})"
    )
    all_w = sorted({col_meta[col]["W"] for col in community_cols}, key=lambda value: int(value[1:]))
    all_c = sorted({int(col_meta[col]["C_num"]) for col in community_cols})
    community_species: dict[int, list[str]] = {}
    community_species_max_relabund: dict[int, dict[str, float]] = {}
    community_species_above_thresh_count: dict[int, dict[str, int]] = {}

    for cnum in all_c:
        c_cols = [col for col in analysis_cols if int(col_meta[col]["C_num"]) == cnum]
        if c_cols:
            sub_rel = candidate_rel_df[c_cols]
            above_count = (sub_rel > RELABUND_MASK).sum(axis=1)
            max_rel = sub_rel.max(axis=1)
        else:
            above_count = pd.Series(0, index=candidate_rel_df.index)
            max_rel = pd.Series(0.0, index=candidate_rel_df.index)
        present_mask = (above_count >= MIN_NONZERO_FOR_PRESENCE) | (
            max_rel >= HIGH_RELABUND_FOR_PRESENCE
        )
        present = present_mask[present_mask].index.tolist()
        present.sort(key=lambda name: int(name.replace("candidate", "")))
        community_species[cnum] = present
        community_species_max_relabund[cnum] = {species: float(max_rel[species]) for species in present}
        community_species_above_thresh_count[cnum] = {
            species: int(above_count[species]) for species in present
        }
        print(f"  Community {cnum}: {len(present)} species")

    present_candidates = sorted(
        {species for species_list in community_species.values() for species in species_list},
        key=lambda name: int(name.replace("candidate", "")),
    )
    present_gids = [candidate_to_gid[species] for species in present_candidates]
    final_species_name = {gid: f"species{idx + 1}" for idx, gid in enumerate(present_gids)}
    old_to_final = {candidate_name[gid]: final_species_name[gid] for gid in present_gids}
    final_species_to_gid = {species: gid for gid, species in final_species_name.items()}

    community_species = {
        cnum: [old_to_final[species] for species in species_list]
        for cnum, species_list in community_species.items()
    }
    community_species_max_relabund = {
        cnum: {
            old_to_final[species]: value
            for species, value in rels.items()
            if species in old_to_final
        }
        for cnum, rels in community_species_max_relabund.items()
    }
    community_species_above_thresh_count = {
        cnum: {
            old_to_final[species]: value
            for species, value in counts.items()
            if species in old_to_final
        }
        for cnum, counts in community_species_above_thresh_count.items()
    }

    final_reads_analysis_df = pd.DataFrame(
        [group_reads_analysis[gid].values for gid in present_gids],
        index=[final_species_name[gid] for gid in present_gids],
        columns=analysis_cols,
    )
    final_reads_analysis_df.index.name = "Species"
    final_reads_all_df = pd.DataFrame(
        [group_reads_all[gid].values for gid in present_gids],
        index=[final_species_name[gid] for gid in present_gids],
        columns=community_cols,
    )
    final_reads_all_df.index.name = "Species"
    print(f"  Final retained species: {len(present_gids)}; renumbered continuously")

    print(f"[{config.title}] [8/8] Export files")
    species_table_rows: list[dict[str, object]] = []
    for rank, gid in enumerate(present_gids, 1):
        rep = group_representative[gid]
        rep_orig = short_to_orig[rep]
        members = group_to_members[gid]
        member_orig = [short_to_orig[member] for member in members]
        rep_taxon = tax_by_id.get(rep_orig, "")
        species_table_rows.append(
            {
                "Species": f"species{rank}",
                "Display_Taxon": "",
                "Representative_ASV": rep,
                "Representative_Original_ID": rep_orig,
                "Genus": genus_or_fallback(rep_taxon),
                "Species_name_from_BLAST": species_name_from_taxon(rep_taxon),
                "Representative_Taxon": rep_taxon,
                "Representative_Taxonomy_confidence": conf_by_id.get(rep_orig, ""),
                "Group_total_reads_analysis": group_total_reads_analysis[gid],
                "Group_total_reads_all_samples": group_total_reads_all[gid],
                "Num_member_ASVs": len(members),
                "Member_ASVs": ",".join(members),
                "Member_Original_IDs": ",".join(member_orig),
            }
        )
    display_by_species = assign_display_taxa(species_table_rows)
    species_table_df = pd.DataFrame(species_table_rows)
    species_table_df.to_csv(paths["species_table"], sep="\t", index=False)

    with paths["species_fasta"].open("w", encoding="utf-8", newline="\n") as handle:
        for rank, gid in enumerate(present_gids, 1):
            rep = group_representative[gid]
            orig = short_to_orig[rep]
            seq = seqs_by_original_id.get(orig)
            if not seq:
                continue
            handle.write(f">species{rank}|FeatureID={orig}|Representative_ASV={rep}\n")
            for idx in range(0, len(seq), 80):
                handle.write(seq[idx : idx + 80] + "\n")

    community_species_rows: list[dict[str, object]] = []
    for cnum in all_c:
        for rank, species in enumerate(community_species[cnum], 1):
            gid = final_species_to_gid[species]
            rep = group_representative[gid]
            rep_orig = short_to_orig[rep]
            above_count = community_species_above_thresh_count[cnum][species]
            rep_taxon = tax_by_id.get(rep_orig, "")
            community_species_rows.append(
                {
                    "Community": cnum,
                    "Rank": rank,
                    "Species": species,
                    "Display_Taxon": display_by_species.get(species, species),
                    "Representative_ASV": rep,
                    "Representative_Original_ID": rep_orig,
                    "Genus": genus_or_fallback(rep_taxon),
                    "Species_name_from_BLAST": species_name_from_taxon(rep_taxon),
                    "Representative_Taxon": rep_taxon,
                    "Max_relabund_in_community": round(
                        community_species_max_relabund[cnum][species], 6
                    ),
                    "Num_samples_above_threshold": above_count,
                    "Presence_rule": "stable"
                    if above_count >= MIN_NONZERO_FOR_PRESENCE
                    else "high_abundance_single_or_pair",
                }
            )
    community_species_df = pd.DataFrame(community_species_rows)
    community_species_df.to_csv(paths["community_species"], sep="\t", index=False)

    species_count_df = pd.DataFrame(
        [{"Community": cnum, "Num_species": len(community_species[cnum])} for cnum in all_c]
    )
    species_count_df.to_csv(paths["species_count"], sep="\t", index=False)

    qc_rows: list[dict[str, object]] = []
    for col in analysis_cols:
        meta = col_meta[col]
        cnum = int(meta["C_num"])
        present = community_species.get(cnum, [])
        kept = float(final_reads_analysis_df.loc[present, col].sum()) if present else 0.0
        total = float(sample_total_reads[col])
        top_total = float(top_total_per_sample[col])
        qc_rows.append(
            {
                "Sample": col,
                "W": meta["W"],
                "Community": cnum,
                "R": meta["R"],
                "D": meta["D"],
                "Total_reads": int(total),
                "TopN_reads": int(top_total),
                "Kept_reads": int(kept),
                "Retention_rate_vs_all_reads": round(kept / total, 6) if total else 0.0,
                "Retention_rate_vs_topN_reads": round(kept / top_total, 6) if top_total else 0.0,
                "Num_species_in_community": len(present),
            }
        )
    qc_df = pd.DataFrame(qc_rows).sort_values(["Community", "W", "R", "D"])
    qc_df.to_csv(paths["retention_qc"], sep="\t", index=False)

    wb = Workbook()
    wb.remove(wb.active)
    for w_value in all_w:
        ws = wb.create_sheet(w_value)
        row = 1
        w_cols = [col for col in community_cols if col_meta[col]["W"] == w_value]
        for cnum in all_c:
            sample_by_slot = {
                (int(col_meta[col]["R_num"]), int(col_meta[col]["D_num"])): col
                for col in w_cols
                if int(col_meta[col]["C_num"]) == cnum
            }
            present = community_species.get(cnum, [])
            ws.cell(row=row, column=1, value=f"Community {cnum}   (n = {len(present)} taxa)")
            row += 1
            for col_idx, header in enumerate(EXCEL_HEADERS, 1):
                ws.cell(row=row, column=col_idx, value=header)
            row += 1
            for species in present:
                ws.cell(row=row, column=1, value=display_by_species.get(species, species))
                for col_idx, slot in enumerate(EXCEL_SLOTS, 2):
                    sample = sample_by_slot.get(slot)
                    value = int(final_reads_all_df.loc[species, sample]) if sample else 1
                    ws.cell(row=row, column=col_idx, value=value)
                row += 1
            row += 2

    for sheet_name, table in [
        ("Community_Species", community_species_df),
        ("Species_Members", species_table_df),
        ("Species_Count", species_count_df),
        ("Retention_QC", qc_df),
        ("ASV_Mapping_Top100", mapping_df.head(100)),
        ("ASV_Pair_Diagnostics", pair_diagnostics_df),
    ]:
        ws = wb.create_sheet(sheet_name)
        write_table_sheet(ws, table)

    readme_ws = wb.create_sheet("README")
    readme_lines = [
        f"{config.title} synthetic-community species identification",
        f"Input feature table: {config.input_table}",
        f"Output directory: {config.output_dir}",
        f"TOP_N_ASV = {config.top_n_asv}",
        f"R_THRESHOLD = {R_THRESHOLD}",
        f"RELABUND_MASK = {RELABUND_MASK}",
        f"MIN_NONZERO_FOR_PRESENCE = {MIN_NONZERO_FOR_PRESENCE}",
        f"HIGH_RELABUND_FOR_PRESENCE = {HIGH_RELABUND_FOR_PRESENCE}",
        "Final species are renumbered continuously after community presence filtering.",
        "Taxon labels are experiment-local ASV-species units: Genus sp.1, Genus sp.2, etc.",
        "Different retained ASV-species groups remain different units even if BLAST names the same species/strain.",
        "Output slots: D1-D10(R1), D8-D10(R2), D8-D10(R3).",
        "Unmeasured plotting slots are filled with 1 for downstream plotting compatibility.",
        "Shared implementation: pipeline/synthetic_community_pipeline.py",
        *config.readme_notes,
    ]
    for idx, line in enumerate(readme_lines, 1):
        readme_ws.cell(row=idx, column=1, value=line)
    readme_ws.column_dimensions["A"].width = 120

    style_excel(wb)
    saved_xlsx = save_workbook(wb, paths["xlsx"])
    paths["xlsx"] = saved_xlsx

    print(f"  Excel: {paths['xlsx']}")
    print(f"  ASV mapping: {paths['asv_mapping']}")
    print(f"  Species table: {paths['species_table']}")
    print(f"  Species fasta: {paths['species_fasta']}")
    print(f"  Community species: {paths['community_species']}")
    print(f"  Species count: {paths['species_count']}")
    print(f"  Retention QC: {paths['retention_qc']}")
    print(f"  ASV pair diagnostics: {paths['pair_diagnostics']}")
    print(
        "  Retention vs topN: "
        f"mean={qc_df['Retention_rate_vs_topN_reads'].mean():.4f}, "
        f"median={qc_df['Retention_rate_vs_topN_reads'].median():.4f}, "
        f"min={qc_df['Retention_rate_vs_topN_reads'].min():.4f}, "
        f"max={qc_df['Retention_rate_vs_topN_reads'].max():.4f}"
    )
    return paths


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the standard synthetic-community pipeline.")
    parser.add_argument(
        "--experiment",
        choices=["all", *EXPERIMENTS.keys()],
        default="all",
        help="Experiment to process.",
    )
    parser.add_argument(
        "--output-dir-name",
        default=None,
        help="Override the output directory name under each experiment root.",
    )
    parser.add_argument(
        "--top-n",
        type=int,
        default=None,
        help="Override TOP_N_ASV for a single experiment run.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.top_n is not None and args.experiment == "all":
        raise ValueError("--top-n can only be used when --experiment is temperature or mortality")
    keys = list(EXPERIMENTS) if args.experiment == "all" else [args.experiment]
    for key in keys:
        config = EXPERIMENTS[key]
        if args.output_dir_name or args.top_n is not None:
            config = ExperimentConfig(
                key=config.key,
                title=config.title,
                root=config.root,
                sample_prefix=config.sample_prefix,
                top_n_asv=args.top_n if args.top_n is not None else config.top_n_asv,
                analysis_filter=config.analysis_filter,
                output_dir_name=args.output_dir_name
                if args.output_dir_name is not None
                else config.output_dir_name,
                taxonomy_file=config.taxonomy_file,
                readme_notes=config.readme_notes,
            )
        run_pipeline(config)


if __name__ == "__main__":
    main()


