"""Standard local BLAST taxonomy annotation for retained synthetic-community species.

The script consumes the representative species FASTA and species table emitted
by ``synthetic_community_pipeline.py``. It writes the same annotation files for
temperature and mortality so downstream code can treat both experiments alike.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from synthetic_community_pipeline import EXPERIMENTS, ExperimentConfig, output_paths


@dataclass
class BlastHit:
    qseqid: str
    sseqid: str
    pident: float
    length: int
    mismatch: int
    gapopen: int
    qstart: int
    qend: int
    qlen: int
    qcovs: float
    evalue: float
    bitscore: float
    stitle: str
    ssciname: str
    sphylum: str
    sclass: str
    sorder: str
    sfamily: str
    sgenus: str


MAPPING_FIELDS = [
    "Experiment",
    "Species",
    "Feature ID",
    "Taxon",
    "Confidence",
    "Representative_ASV",
    "Num_member_ASVs",
    "Member_Original_IDs",
    "Query_length",
    "Status",
    "Identity_pct",
    "Coverage_pct",
    "Alignment_length",
    "Evalue",
    "Bitscore",
    "Top_accession",
    "Scientific_name",
    "Top_hit_description",
    "Annotation_source",
]
TAXONOMY_FIELDS = ["Feature ID", "Taxon", "Confidence"]
HIT_FIELDS = [
    "Species",
    "Rank",
    "Feature ID",
    "Accession",
    "Identity_pct",
    "Coverage_pct",
    "Alignment_length",
    "Evalue",
    "Bitscore",
    "Scientific_name",
    "Taxon",
    "Top_hit_description",
]


def standard_annotation_dir(config: ExperimentConfig) -> Path:
    return config.root / "species_annotations" / "local_blast_standard"


def annotation_paths(config: ExperimentConfig, out_dir: Path | None = None) -> dict[str, Path]:
    out = out_dir or standard_annotation_dir(config)
    label = config.label
    return {
        "raw": out / f"{label}_blast_hits_raw.tsv",
        "mapping": out / f"{label}_species_taxonomy_mapping.tsv",
        "taxonomy": out / f"{label}_taxonomy_style.tsv",
        "hits": out / f"{label}_all_blast_hits_ranked.tsv",
        "xlsx": out / f"{label}_species_taxonomy.xlsx",
    }


def iter_fasta(path: Path):
    name = None
    parts: list[str] = []
    with path.open("r", encoding="utf-8") as handle:
        for raw in handle:
            line = raw.strip()
            if not line:
                continue
            if line.startswith(">"):
                if name is not None:
                    yield name, "".join(parts)
                name = line[1:].split()[0]
                parts = []
            else:
                parts.append(line)
    if name is not None:
        yield name, "".join(parts)


def read_species_table(path: Path) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    return {row["Species"]: row for row in rows}


def read_species_meta_from_query(query_fasta: Path) -> dict[str, dict[str, str]]:
    meta: dict[str, dict[str, str]] = {}
    for record_id, _seq in iter_fasta(query_fasta):
        parts = record_id.split("|")
        species = parts[0]
        values = {
            "Species": species,
            "Representative_Original_ID": "",
            "Representative_ASV": "",
            "Num_member_ASVs": "",
            "Member_Original_IDs": "",
        }
        for part in parts[1:]:
            if "=" not in part:
                continue
            key, value = part.split("=", 1)
            if key == "FeatureID":
                values["Representative_Original_ID"] = value
            elif key == "Representative_ASV":
                values["Representative_ASV"] = value
        meta[species] = values
    return meta


def clean_tax(value: object) -> str:
    text = str(value or "").strip()
    return "" if text in {"N/A", "NA", "nan", "None"} else text


def split_species_name(scientific_name: str) -> tuple[str, str]:
    name = clean_tax(scientific_name)
    if not name:
        return "", ""
    pieces = name.split()
    if len(pieces) >= 2 and pieces[0][0].isupper():
        return pieces[0], " ".join(pieces[:2])
    return pieces[0] if pieces else "", name


def confidence_for_16s(identity: float, coverage: float, aln_len: int) -> str:
    if identity >= 99.0 and coverage >= 90 and aln_len >= 250:
        return "high_genus_possible_species"
    if identity >= 97.0 and coverage >= 85:
        return "medium_genus"
    if identity >= 95.0 and coverage >= 80:
        return "low_family_or_genus"
    if identity >= 90.0 and coverage >= 70:
        return "very_low"
    return "unreliable"


def taxonomy_string(hit: BlastHit) -> str:
    genus_from_name, species_from_name = split_species_name(hit.ssciname)
    genus = clean_tax(hit.sgenus) or genus_from_name
    levels = [
        ("d", "Bacteria"),
        ("p", hit.sphylum),
        ("c", hit.sclass),
        ("o", hit.sorder),
        ("f", hit.sfamily),
        ("g", genus),
    ]
    taxon = "; ".join(
        f"{prefix}__{clean_tax(value)}" for prefix, value in levels if clean_tax(value)
    )
    if species_from_name and genus and species_from_name.startswith(genus + " "):
        taxon = f"{taxon}; s__{species_from_name}" if taxon else f"s__{species_from_name}"
    return taxon


def run_blast(
    query: Path,
    raw_out: Path,
    blast_db: str,
    blastdb_dir: Path | None,
    threads: int,
    max_target_seqs: int,
    min_identity: float,
    min_qcov: float,
) -> None:
    if shutil.which("blastn") is None:
        raise RuntimeError(
            "blastn was not found on PATH. Install BLAST+ or rerun with --skip-blast "
            "after placing an existing raw BLAST output at the expected raw TSV path."
        )

    env = os.environ.copy()
    if blastdb_dir is not None:
        existing = env.get("BLASTDB", "")
        env["BLASTDB"] = str(blastdb_dir) if not existing else f"{blastdb_dir}{os.pathsep}{existing}"

    outfmt = (
        "6 qseqid sseqid pident length mismatch gapopen qstart qend qlen qcovs "
        "evalue bitscore stitle ssciname sphylum sclass sorder sfamily sgenus"
    )
    cmd = [
        "blastn",
        "-query",
        str(query),
        "-db",
        blast_db,
        "-outfmt",
        outfmt,
        "-max_target_seqs",
        str(max_target_seqs),
        "-perc_identity",
        str(min_identity),
        "-qcov_hsp_perc",
        str(min_qcov),
        "-num_threads",
        str(threads),
    ]
    with raw_out.open("w", encoding="utf-8", newline="") as handle:
        subprocess.run(cmd, check=True, stdout=handle, env=env)


def parse_blast(raw_out: Path) -> dict[str, list[BlastHit]]:
    cols = [
        "qseqid",
        "sseqid",
        "pident",
        "length",
        "mismatch",
        "gapopen",
        "qstart",
        "qend",
        "qlen",
        "qcovs",
        "evalue",
        "bitscore",
        "stitle",
        "ssciname",
        "sphylum",
        "sclass",
        "sorder",
        "sfamily",
        "sgenus",
    ]
    if not raw_out.exists() or raw_out.stat().st_size == 0:
        return {}
    df = pd.read_csv(raw_out, sep="\t", names=cols, keep_default_na=False)
    hits: dict[str, list[BlastHit]] = {}
    for row in df.itertuples(index=False):
        hit = BlastHit(
            qseqid=str(row.qseqid),
            sseqid=str(row.sseqid),
            pident=float(row.pident),
            length=int(row.length),
            mismatch=int(row.mismatch),
            gapopen=int(row.gapopen),
            qstart=int(row.qstart),
            qend=int(row.qend),
            qlen=int(row.qlen),
            qcovs=float(row.qcovs),
            evalue=float(row.evalue),
            bitscore=float(row.bitscore),
            stitle=str(row.stitle),
            ssciname=str(row.ssciname),
            sphylum=str(row.sphylum),
            sclass=str(row.sclass),
            sorder=str(row.sorder),
            sfamily=str(row.sfamily),
            sgenus=str(row.sgenus),
        )
        species = hit.qseqid.split("|", 1)[0]
        hits.setdefault(species, []).append(hit)
    for species in hits:
        hits[species].sort(key=lambda item: (-item.bitscore, -item.pident, -item.qcovs, item.evalue))
    return hits


def query_lengths(query_fasta: Path) -> dict[str, int]:
    lengths: dict[str, int] = {}
    for record_id, seq in iter_fasta(query_fasta):
        species = record_id.split("|", 1)[0]
        lengths[species] = len(seq)
    return lengths


def species_sort_key(species: str) -> tuple[int, str]:
    match = re.search(r"(\d+)$", species)
    return (int(match.group(1)) if match else 10**9, species)


def build_tables(
    experiment_key: str,
    species_meta: dict[str, dict[str, str]],
    hits_by_species: dict[str, list[BlastHit]],
    lengths: dict[str, int],
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    mapping_rows: list[dict[str, object]] = []
    taxonomy_rows: list[dict[str, object]] = []
    all_hit_rows: list[dict[str, object]] = []

    for species in sorted(species_meta, key=species_sort_key):
        meta = species_meta[species]
        hits = hits_by_species.get(species, [])
        best = hits[0] if hits else None
        feature_id = meta.get("Representative_Original_ID", "")
        if best is None:
            taxon = ""
            confidence = ""
            row = {
                "Experiment": experiment_key,
                "Species": species,
                "Feature ID": feature_id,
                "Taxon": taxon,
                "Confidence": confidence,
                "Representative_ASV": meta.get("Representative_ASV", ""),
                "Num_member_ASVs": meta.get("Num_member_ASVs", ""),
                "Member_Original_IDs": meta.get("Member_Original_IDs", ""),
                "Query_length": lengths.get(species, ""),
                "Status": "no_hit",
                "Identity_pct": "",
                "Coverage_pct": "",
                "Alignment_length": "",
                "Evalue": "",
                "Bitscore": "",
                "Top_accession": "",
                "Scientific_name": "",
                "Top_hit_description": "",
                "Annotation_source": "local_blast",
            }
        else:
            taxon = taxonomy_string(best)
            confidence = confidence_for_16s(best.pident, best.qcovs, best.length)
            row = {
                "Experiment": experiment_key,
                "Species": species,
                "Feature ID": feature_id,
                "Taxon": taxon,
                "Confidence": confidence,
                "Representative_ASV": meta.get("Representative_ASV", ""),
                "Num_member_ASVs": meta.get("Num_member_ASVs", ""),
                "Member_Original_IDs": meta.get("Member_Original_IDs", ""),
                "Query_length": lengths.get(species, ""),
                "Status": "success",
                "Identity_pct": best.pident,
                "Coverage_pct": best.qcovs,
                "Alignment_length": best.length,
                "Evalue": best.evalue,
                "Bitscore": best.bitscore,
                "Top_accession": best.sseqid,
                "Scientific_name": clean_tax(best.ssciname),
                "Top_hit_description": best.stitle,
                "Annotation_source": "local_blast",
            }
        mapping_rows.append(row)
        taxonomy_rows.append({"Feature ID": feature_id, "Taxon": taxon, "Confidence": confidence})

        for rank, hit in enumerate(hits, 1):
            all_hit_rows.append(
                {
                    "Species": species,
                    "Rank": rank,
                    "Feature ID": feature_id,
                    "Accession": hit.sseqid,
                    "Identity_pct": hit.pident,
                    "Coverage_pct": hit.qcovs,
                    "Alignment_length": hit.length,
                    "Evalue": hit.evalue,
                    "Bitscore": hit.bitscore,
                    "Scientific_name": clean_tax(hit.ssciname),
                    "Taxon": taxonomy_string(hit),
                    "Top_hit_description": hit.stitle,
                }
            )

    return mapping_rows, taxonomy_rows, all_hit_rows


def write_tsv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def style_header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_excel(path: Path, sheets: dict[str, tuple[list[dict[str, object]], list[str]]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    confidence_fill = {
        "high_genus_possible_species": "D5F5E3",
        "medium_genus": "FEF9E7",
        "low_family_or_genus": "FDEBD0",
        "very_low": "FDEDEC",
        "unreliable": "FADBD8",
    }

    for sheet_name, (rows, fields) in sheets.items():
        ws = wb.create_sheet(sheet_name)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for col_idx, field in enumerate(fields, 1):
            cell = ws.cell(row=1, column=col_idx, value=field)
            style_header(cell)
            cell.border = border
        for row_idx, row in enumerate(rows, 2):
            for col_idx, field in enumerate(fields, 1):
                value = row.get(field, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=field in {"Taxon", "Top_hit_description", "Member_Original_IDs"},
                )
                if field == "Confidence" and str(value) in confidence_fill:
                    cell.fill = PatternFill("solid", fgColor=confidence_fill[str(value)])
        for col_idx, field in enumerate(fields, 1):
            width = 14
            if field in {"Taxon", "Top_hit_description"}:
                width = 70
            elif field in {"Feature ID", "Member_Original_IDs"}:
                width = 36
            elif field in {"Species", "Representative_ASV", "Confidence"}:
                width = 22
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    readme = wb.create_sheet("README")
    lines = [
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "Input: representative retained species sequences from the standard ASV pipeline.",
        "Taxonomy source: local BLAST database selected by --blast-db.",
        "The same schema is used for temperature and mortality.",
    ]
    for idx, line in enumerate(lines, 1):
        readme.cell(row=idx, column=1, value=line)
    readme.column_dimensions["A"].width = 110
    try:
        wb.save(path)
        return path
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_updated{path.suffix}")
        wb.save(fallback)
        return fallback


def annotate_experiment(args: argparse.Namespace, config: ExperimentConfig) -> dict[str, Path]:
    pipeline_paths = output_paths(config)
    query = Path(args.query) if args.query else pipeline_paths["species_fasta"]
    species_table = Path(args.species_table) if args.species_table else pipeline_paths["species_table"]
    out_dir = Path(args.out_dir) if args.out_dir else standard_annotation_dir(config)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = annotation_paths(config, out_dir)

    if not query.exists():
        raise FileNotFoundError(f"Query FASTA not found: {query}")
    if args.query and args.experiment == "all":
        raise ValueError("--query can only be used with a single --experiment")
    if args.species_table and args.experiment == "all":
        raise ValueError("--species-table can only be used with a single --experiment")
    if args.out_dir and args.experiment == "all":
        raise ValueError("--out-dir can only be used with a single --experiment")

    print(f"[{config.title}] Query FASTA: {query}")
    print(f"[{config.title}] Species table: {species_table}")
    print(f"[{config.title}] BLAST DB: {args.blast_db}")
    if args.blastdb_dir is not None:
        print(f"[{config.title}] BLASTDB dir: {args.blastdb_dir}")

    if not args.skip_blast:
        run_blast(
            query=query,
            raw_out=paths["raw"],
            blast_db=args.blast_db,
            blastdb_dir=args.blastdb_dir,
            threads=args.threads,
            max_target_seqs=args.max_target_seqs,
            min_identity=args.min_identity,
            min_qcov=args.min_qcov,
        )
    elif not paths["raw"].exists():
        raise FileNotFoundError(f"--skip-blast was used, but raw BLAST TSV does not exist: {paths['raw']}")

    species_meta = (
        read_species_table(species_table)
        if species_table.exists()
        else read_species_meta_from_query(query)
    )
    hits_by_species = parse_blast(paths["raw"])
    lengths = query_lengths(query)
    mapping_rows, taxonomy_rows, all_hit_rows = build_tables(
        config.key, species_meta, hits_by_species, lengths
    )

    write_tsv(paths["mapping"], mapping_rows, MAPPING_FIELDS)
    write_tsv(paths["taxonomy"], taxonomy_rows, TAXONOMY_FIELDS)
    write_tsv(paths["hits"], all_hit_rows, HIT_FIELDS)
    paths["xlsx"] = write_excel(
        paths["xlsx"],
        {
            "Species_mapping": (mapping_rows, MAPPING_FIELDS),
            "taxonomy.tsv_style": (taxonomy_rows, TAXONOMY_FIELDS),
            "All_BLAST_hits": (all_hit_rows, HIT_FIELDS),
        },
    )

    n_success = sum(1 for row in mapping_rows if row["Status"] == "success")
    print(f"[{config.title}] Annotated: {n_success}/{len(mapping_rows)} species")
    print(f"[{config.title}] Mapping TSV: {paths['mapping']}")
    print(f"[{config.title}] taxonomy.tsv style TSV: {paths['taxonomy']}")
    print(f"[{config.title}] Excel: {paths['xlsx']}")
    return paths


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Annotate retained species with local BLAST.")
    parser.add_argument("--experiment", choices=["all", *EXPERIMENTS.keys()], default="all")
    parser.add_argument("--query", type=Path, default=None)
    parser.add_argument("--species-table", type=Path, default=None)
    parser.add_argument("--out-dir", type=Path, default=None)
    parser.add_argument("--blast-db", default="16S_ribosomal_RNA")
    parser.add_argument("--blastdb-dir", type=Path, default=None)
    parser.add_argument("--threads", type=int, default=8)
    parser.add_argument("--max-target-seqs", type=int, default=10)
    parser.add_argument("--min-identity", type=float, default=90.0)
    parser.add_argument("--min-qcov", type=float, default=70.0)
    parser.add_argument(
        "--skip-blast",
        action="store_true",
        help="Parse an existing raw BLAST TSV instead of running blastn.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    keys = list(EXPERIMENTS) if args.experiment == "all" else [args.experiment]
    for key in keys:
        annotate_experiment(args, EXPERIMENTS[key])


if __name__ == "__main__":
    main()
